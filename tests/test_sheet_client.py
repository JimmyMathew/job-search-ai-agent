import json

import config
import openpyxl
import pytest
import sheet
from config import UserFacingError
from openpyxl.worksheet.datavalidation import DataValidation
from sheet_schema import (
    JOBS_HEADERS,
    JOBS_TAB,
    LEGACY_STATUS_SNAPSHOT_KEY,
    META_TAB,
    PREFERENCES_TAB,
    RULES_TAB,
    STATUS_SYNC_TAB,
    TAB_HEADERS,
    WHY_PENDING,
)
from util import run_cli
from workbook import (
    LOCKED_PREFIX,
    WorkbookCorrupt,
    WorkbookError,
    WorkbookLocked,
    WorkbookMissing,
)


@pytest.fixture
def profile_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "PROFILES_DIR", tmp_path)
    folder = tmp_path / "jane"
    folder.mkdir()
    (folder / "config.yaml").write_text('locations: ["Berlin"]\n', encoding="utf-8")
    return folder


@pytest.fixture
def client(profile_dir):
    client = sheet.SheetClient("jane")
    client.init()
    return client


def make_job(job_id, **extra):
    return {
        "job_id": str(job_id),
        "title": "Analyst",
        "company": "Acme",
        "location": "Berlin",
        "posted_date": "2026-10-01",
        "link": f"https://www.linkedin.com/jobs/view/{job_id}",
        **extra,
    }


def edit_cells(client, tab, row, **values):
    """What the person does in Excel: type into cells of an existing row, then save."""
    book = openpyxl.load_workbook(client.workbook.path)
    ws = book[tab]
    columns = {cell.value: cell.column for cell in ws[1]}
    for header, value in values.items():
        ws.cell(row, columns[header]).value = value
    book.save(client.workbook.path)


def jobs_by_id(client):
    return {r["Job ID"]: r for r in client.workbook.get_records(JOBS_TAB)}


class TestInit:
    def test_creates_the_workbook_beside_the_profile_with_every_tab(self, profile_dir, client):
        assert client.workbook.path == profile_dir / "Job Search - jane.xlsx"
        book = openpyxl.load_workbook(client.workbook.path)
        assert book.sheetnames == list(TAB_HEADERS)
        assert book.active.title == JOBS_TAB
        assert [c.value for c in book[JOBS_TAB][1]] == JOBS_HEADERS

    def test_bookkeeping_tabs_are_hidden_and_the_rest_are_not(self, client):
        book = openpyxl.load_workbook(client.workbook.path)
        hidden = {ws.title for ws in book if ws.sheet_state == "hidden"}
        assert hidden == {META_TAB, STATUS_SYNC_TAB}

    def test_the_status_dropdown_is_on_the_status_column(self, client):
        ws = openpyxl.load_workbook(client.workbook.path)[JOBS_TAB]
        (validation,) = ws.data_validations.dataValidation
        letter = openpyxl.utils.get_column_letter(JOBS_HEADERS.index("Status") + 1)
        assert str(validation.sqref).startswith(f"{letter}2")
        assert "Interview" in validation.formula1

    def test_records_the_workbook_in_the_state_file(self, profile_dir, client):
        state = config.read_state("jane")
        assert state["workbook"] == str(profile_dir / "Job Search - jane.xlsx")
        assert state["created_at"]

    def test_running_again_changes_nothing_and_keeps_the_creation_time_and_other_state(
        self, client
    ):
        config.write_state("jane", {**config.read_state("jane"), "note": "keep me"})
        created_at = config.read_state("jane")["created_at"]
        before = client.workbook.path.read_bytes()

        sheet.SheetClient("jane").init()

        assert client.workbook.path.read_bytes() == before
        assert config.read_state("jane")["created_at"] == created_at
        assert config.read_state("jane")["note"] == "keep me"

    def test_an_existing_workbook_is_extended_never_overwritten(self, profile_dir, client):
        client.append_jobs([make_job(1)])
        edit_cells(client, PREFERENCES_TAB, 3, Value="60,000 a year")
        book = openpyxl.load_workbook(client.workbook.path)
        del book[RULES_TAB]
        book.save(client.workbook.path)
        (profile_dir / "agent-state.json").unlink()  # e.g. the state file was lost

        sheet.SheetClient("jane").init()

        book = openpyxl.load_workbook(client.workbook.path)
        assert RULES_TAB in book.sheetnames
        assert list(jobs_by_id(client)) == ["1"]
        (answer,) = client.get_preferences()
        assert answer["value"] == "60,000 a year"

    def test_a_workbook_the_state_knows_but_that_is_gone_is_explained_not_recreated(self, client):
        client.workbook.path.unlink()

        with pytest.raises(WorkbookMissing) as error:
            sheet.SheetClient("jane").init()

        message = str(error.value)
        assert str(client.workbook.path) in message
        assert str(client.workbook.backups_dir) in message
        assert "agent-state.json" in message
        assert "forgets" in message
        assert not client.workbook.path.exists()

    def test_the_workbook_setting_moves_the_file(self, profile_dir):
        (profile_dir / "config.yaml").write_text(
            'locations: ["Berlin"]\nworkbook: "lists/mine.xlsx"\n', encoding="utf-8"
        )
        moved = sheet.SheetClient("jane")
        moved.init()
        assert moved.workbook.path == profile_dir / "lists" / "mine.xlsx"
        assert moved.workbook.path.exists()

    def test_every_other_command_says_to_run_init_when_there_is_no_list(self, profile_dir):
        with pytest.raises(WorkbookMissing, match="sheet.py --profile jane init"):
            sheet.SheetClient("jane").get_meta()


class TestReady:
    def test_checks_the_file_and_takes_one_backup_per_call(self, client):
        assert client.ready() == client.workbook.path
        assert len(client.workbook.backup_files()) == 1

    def test_an_open_workbook_is_reported(self, client):
        client.workbook.path.with_name(f"~${client.workbook.path.name}").write_bytes(b"")
        with pytest.raises(WorkbookLocked, match=LOCKED_PREFIX):
            client.ready()
        assert client.workbook.backup_files() == []

    def test_a_damaged_file_is_reported_and_not_backed_up(self, client):
        client.ready()
        client.workbook.path.write_bytes(b"damaged")
        with pytest.raises(WorkbookCorrupt):
            client.ready()
        assert len(client.workbook.backup_files()) == 1

    def test_a_missing_file_is_reported(self, client):
        client.workbook.path.unlink()
        with pytest.raises(WorkbookMissing):
            client.ready()


class TestDedupe:
    def test_jobs_already_listed_are_dropped_and_new_ones_kept(self, client):
        client.append_jobs([make_job(1)])
        result = client.dedupe([make_job(1), make_job(2, title="Other")])
        assert [j["job_id"] for j in result["new"]] == ["2"]
        assert result["reposts"] == []

    def test_same_title_and_company_near_a_listed_one_is_a_repost(self, client):
        client.append_jobs([make_job(1)])
        result = client.dedupe([make_job(2, posted_date="2026-10-10")])
        assert [j["job_id"] for j in result["reposts"]] == ["2"]
        assert result["new"] == []

    def test_a_repost_outside_the_window_is_new(self, client):
        client.append_jobs([make_job(1)])
        result = client.dedupe([make_job(2, posted_date="2026-12-31")])
        assert [j["job_id"] for j in result["new"]] == ["2"]

    def test_an_empty_list_dedupes_everything_as_new(self, client):
        assert len(client.dedupe([make_job(1), make_job(2)])["new"]) == 2


class TestAppendJobs:
    def test_rows_land_in_the_right_columns_with_defaults(self, client):
        added = client.append_jobs([make_job(1, why="Good fit", score=82, flags=["a", "b"])])

        assert added == 1
        row = jobs_by_id(client)["1"]
        assert row["Job title"] == "Analyst"
        assert row["Company"] == "Acme"
        assert row["Score"] == "82"
        assert row["Why it matches"] == "Good fit"
        assert row["Flags"] == "a, b"
        assert row["Status"] == "New"
        assert row["Reason"] == ""
        assert len(row["Date added"]) == len("2026-10-12")

    def test_score_is_stored_as_a_number_and_the_job_id_as_text(self, client):
        client.append_jobs([make_job(3912345678, score=82)])
        ws = openpyxl.load_workbook(client.workbook.path)[JOBS_TAB]
        assert ws["B2"].value == 82
        assert ws["C2"].value == "3912345678"
        assert ws["C2"].data_type == "s"

    def test_retrying_the_same_write_adds_nothing(self, client):
        jobs = [make_job(1), make_job(2)]
        assert client.append_jobs(jobs) == 2
        assert client.append_jobs(jobs) == 0
        assert list(jobs_by_id(client)) == ["1", "2"]

    def test_a_reordered_file_is_written_by_header_name(self, client):
        book = openpyxl.load_workbook(client.workbook.path)
        ws = book[JOBS_TAB]
        ws["A1"], ws["D1"] = "Job title", "Date added"
        book.save(client.workbook.path)

        client.append_jobs([make_job(1)])

        assert jobs_by_id(client)["1"]["Job title"] == "Analyst"

    def test_text_that_looks_like_a_formula_stays_text(self, client):
        client.append_jobs([make_job(1, title="=SUM(1,1)")])
        assert jobs_by_id(client)["1"]["Job title"] == "=SUM(1,1)"


class TestUpdateJobs:
    def test_only_scoring_columns_change(self, client):
        client.append_jobs([make_job(1, why=WHY_PENDING)])
        edit_cells(client, JOBS_TAB, 2, Status="Applied", Reason="Called them")

        updated = client.update_jobs(
            [
                make_job(
                    1,
                    score=77,
                    why="Strong match",
                    gaps="No SQL",
                    flags=["onsite"],
                    role_type="target",
                    title="CHANGED",
                    filtered_reason="ignored",
                )
            ]
        )

        row = jobs_by_id(client)["1"]
        assert updated == 1
        assert (row["Score"], row["Why it matches"], row["Gaps"]) == (
            "77",
            "Strong match",
            "No SQL",
        )
        assert (row["Flags"], row["Role type"]) == ("onsite", "target")
        assert (row["Status"], row["Reason"]) == ("Applied", "Called them")
        assert row["Job title"] == "Analyst"

    def test_unknown_jobs_are_skipped_never_appended(self, client):
        client.append_jobs([make_job(1)])
        assert client.update_jobs([{"job_id": "99", "score": 5, "why": "x"}]) == 0
        assert list(jobs_by_id(client)) == ["1"]

    def test_a_job_id_given_as_a_number_still_matches(self, client):
        client.append_jobs([make_job(1)])
        assert client.update_jobs([{"job_id": 1, "score": 5, "why": "x"}]) == 1

    def test_a_sorted_list_is_updated_on_the_right_rows(self, client):
        client.append_jobs([make_job(i, title=f"Job {i}") for i in (1, 2, 3)])
        book = openpyxl.load_workbook(client.workbook.path)
        ws = book[JOBS_TAB]
        for row in ws.iter_rows(min_row=2, max_row=4):
            row[2].value = {"1": "3", "2": "2", "3": "1"}[row[2].value]  # reverse the Job IDs
        for row, title in zip(
            ws.iter_rows(min_row=2, max_row=4), ("Job 3", "Job 2", "Job 1"), strict=True
        ):
            row[3].value = title
        book.save(client.workbook.path)

        client.update_jobs([{"job_id": "1", "score": 91, "why": "ok"}])

        assert jobs_by_id(client)["1"]["Score"] == "91"
        assert jobs_by_id(client)["1"]["Job title"] == "Job 1"
        assert jobs_by_id(client)["3"]["Score"] == ""

    def test_a_locked_file_changes_nothing(self, client):
        client.append_jobs([make_job(1)])
        client.workbook.path.with_name(f"~${client.workbook.path.name}").write_bytes(b"")
        before = client.workbook.path.read_bytes()
        with pytest.raises(WorkbookLocked):
            client.update_jobs([{"job_id": "1", "score": 5, "why": "x"}])
        assert client.workbook.path.read_bytes() == before


class TestPending:
    def test_pending_and_failed_rows_are_returned_as_job_dicts(self, client):
        client.append_jobs(
            [
                make_job(1, why=WHY_PENDING),
                make_job(2, why="Scoring failed — will retry next run"),
                make_job(3, why="Not scored this run (over the per-run limit)"),
                make_job(4, why="Blocked by rule: x"),
                make_job(5, why="fine", score=70),
            ]
        )

        pending = client.get_pending_jobs()

        assert [j["job_id"] for j in pending] == ["1", "2"]
        assert pending[0] == {
            "job_id": "1",
            "title": "Analyst",
            "company": "Acme",
            "location": "Berlin",
            "posted_date": "2026-10-01",
            "link": "https://www.linkedin.com/jobs/view/1",
        }

    def test_a_scored_row_is_no_longer_pending(self, client):
        client.append_jobs([make_job(1, why=WHY_PENDING)])
        client.update_jobs([{"job_id": "1", "score": 60, "why": "ok"}])
        assert client.get_pending_jobs() == []


class TestFeedback:
    def read_and_mark(self, client):
        changed = client.get_changed_status_rows()
        client.mark_status_synced(changed)
        return changed

    def test_fresh_rows_are_not_feedback(self, client):
        client.append_jobs([make_job(1)])
        assert client.get_changed_status_rows() == []

    def test_a_changed_status_is_reported_until_it_is_marked_as_read(self, client):
        client.append_jobs([make_job(1), make_job(2)])
        edit_cells(client, JOBS_TAB, 2, Status="Not for me", Reason="too junior")

        (changed,) = client.get_changed_status_rows()
        assert changed == {
            "job_id": "1",
            "title": "Analyst",
            "company": "Acme",
            "status": "Not for me",
            "reason": "too junior",
        }

        client.mark_status_synced([changed])
        assert client.get_changed_status_rows() == []

        edit_cells(client, JOBS_TAB, 2, Reason="wrong industry")
        assert [c["reason"] for c in client.get_changed_status_rows()] == ["wrong industry"]

    def test_an_edit_made_after_the_rows_were_read_is_still_reported_next_time(self, client):
        client.append_jobs([make_job(1)])
        edit_cells(client, JOBS_TAB, 2, Status="Not for me", Reason="too junior")
        changed = client.get_changed_status_rows()

        edit_cells(client, JOBS_TAB, 2, Reason="actually, wrong industry")  # between read and mark
        client.mark_status_synced(changed)

        assert client.workbook.get_records(STATUS_SYNC_TAB) == [
            {"Job ID": "1", "Status": "Not for me", "Reason": "too junior"}
        ]
        (again,) = client.get_changed_status_rows()
        assert again["reason"] == "actually, wrong industry"

    def test_marking_twice_updates_the_same_row(self, client):
        client.append_jobs([make_job(1)])
        edit_cells(client, JOBS_TAB, 2, Status="Applied")
        self.read_and_mark(client)
        edit_cells(client, JOBS_TAB, 2, Status="Interview")
        self.read_and_mark(client)

        assert client.workbook.get_records(STATUS_SYNC_TAB) == [
            {"Job ID": "1", "Status": "Interview", "Reason": ""}
        ]

    def test_marking_nothing_changes_nothing(self, client):
        before = client.workbook.path.read_bytes()
        client.mark_status_synced([])
        assert client.workbook.path.read_bytes() == before

    def test_reverting_to_new_after_feedback_is_reported(self, client):
        client.append_jobs([make_job(1)])
        edit_cells(client, JOBS_TAB, 2, Status="Not for me")
        self.read_and_mark(client)
        edit_cells(client, JOBS_TAB, 2, Status="New")
        assert [c["status"] for c in client.get_changed_status_rows()] == ["New"]

    @pytest.mark.parametrize("bad", [{"job_id": "1"}, {"status": "Applied"}, "nonsense"])
    def test_a_malformed_row_is_a_plain_error_and_nothing_is_written(self, client, bad):
        before = client.workbook.path.read_bytes()
        with pytest.raises(UserFacingError, match="changed row #1"):
            client.mark_status_synced([bad])
        assert client.workbook.path.read_bytes() == before

    def test_the_file_must_be_a_list(self, client):
        with pytest.raises(UserFacingError, match="JSON list"):
            client.mark_status_synced({"job_id": "1"})


class TestRules:
    def rule(self, **extra):
        return {"rule": "Avoid sales", "type": "Block", "keywords": "sales;quota", **extra}

    def test_a_learned_rule_is_marked_as_written_by_the_agent_and_active(self, client):
        client.append_rules([self.rule(based_on="jobs 1 and 2")])

        assert client.get_rules() == [
            {
                "rule": "Avoid sales",
                "type": "Block",
                "keywords": "sales;quota",
                "based_on": "jobs 1 and 2",
                "added_by": sheet.AGENT_RULE_AUTHOR,
                "active": "Yes",
            }
        ]
        assert sheet.AGENT_RULE_AUTHOR == "AI"

    def test_adding_the_same_rule_again_does_not_duplicate_it(self, client):
        client.append_rules([self.rule()])
        client.append_rules([self.rule(), self.rule(rule="Prefer remote", type="Boost")])
        assert [r["rule"] for r in client.get_rules()] == ["Avoid sales", "Prefer remote"]

    def test_a_rule_written_by_the_person_is_kept_as_theirs(self, client):
        client.append_rules([self.rule(added_by="Jane", active="No")])
        (rule,) = client.get_rules()
        assert (rule["added_by"], rule["active"]) == ("Jane", "No")

    def test_no_rules_by_default(self, client):
        assert client.get_rules() == []


class TestPreferences:
    def test_seeded_with_questions_and_nothing_returned_until_answered(self, client):
        assert len(client.workbook.get_records(PREFERENCES_TAB)) > 5
        assert client.get_preferences() == []

    def test_only_answered_rows_are_returned(self, client):
        edit_cells(client, PREFERENCES_TAB, 3, Value="60,000 per year")
        edit_cells(client, PREFERENCES_TAB, 4, Value="   ")

        (answered,) = client.get_preferences()

        assert answered["value"] == "60,000 per year"
        assert "salary" in answered["criterion"].lower()

    def test_the_answers_survive_running_init_again(self, client):
        edit_cells(client, PREFERENCES_TAB, 3, Value="60,000 per year")
        client.init()
        assert [p["value"] for p in client.get_preferences()] == ["60,000 per year"]


class TestMeta:
    def test_new_keys_are_added_and_existing_ones_updated_in_place(self, client):
        client.update_meta({"last_run_iso": "2026-10-01T10:00:00+00:00", "jobs_seen_last_run": "5"})
        client.update_meta({"jobs_seen_last_run": "9", "linkedin_status": "ok"})

        assert client.get_meta() == {
            "last_run_iso": "2026-10-01T10:00:00+00:00",
            "jobs_seen_last_run": "9",
            "linkedin_status": "ok",
        }
        assert len(client.workbook.get_records(META_TAB)) == 3

    def test_empty_meta(self, client):
        assert client.get_meta() == {}


class TestCommandLine:
    def run(self, monkeypatch, *argv):
        monkeypatch.setattr("sys.argv", ["sheet.py", "--profile", "jane", *argv])
        run_cli(sheet.main)

    def test_init_prints_the_workbook_path_and_init_tabs_is_an_alias(
        self, profile_dir, monkeypatch, capsys
    ):
        self.run(monkeypatch, "init")
        self.run(monkeypatch, "init-tabs")
        lines = capsys.readouterr().out.splitlines()
        assert lines == [f"WORKBOOK: {profile_dir / 'Job Search - jane.xlsx'}"] * 2

    def test_where_works_before_the_file_exists(self, profile_dir, monkeypatch, capsys):
        self.run(monkeypatch, "where")
        assert (
            capsys.readouterr().out.strip() == f"WORKBOOK: {profile_dir / 'Job Search - jane.xlsx'}"
        )

    def test_ready_prints_the_path(self, client, monkeypatch, capsys):
        self.run(monkeypatch, "ready")
        assert capsys.readouterr().out.strip() == f"WORKBOOK: {client.workbook.path}"

    def test_a_workbook_open_in_excel_ends_a_command_with_the_message_to_show(
        self, client, tmp_path, monkeypatch
    ):
        client.workbook.path.with_name(f"~${client.workbook.path.name}").write_bytes(b"")
        jobs_file = tmp_path / "jobs.json"
        jobs_file.write_text(json.dumps([make_job(1)]), encoding="utf-8")

        with pytest.raises(SystemExit) as exit_info:
            self.run(monkeypatch, "append-jobs", "--in", str(jobs_file))

        assert str(exit_info.value).startswith(LOCKED_PREFIX)

    def test_the_same_command_works_once_it_is_closed(self, client, tmp_path, monkeypatch, capsys):
        lock = client.workbook.path.with_name(f"~${client.workbook.path.name}")
        lock.write_bytes(b"")
        jobs_file = tmp_path / "jobs.json"
        jobs_file.write_text(json.dumps({"jobs": [make_job(1)]}), encoding="utf-8")
        with pytest.raises(SystemExit):
            self.run(monkeypatch, "append-jobs", "--in", str(jobs_file))

        lock.unlink()
        self.run(monkeypatch, "append-jobs", "--in", str(jobs_file))

        assert "Appended 1 rows to Jobs (0 already present)" in capsys.readouterr().out

    def test_clear_lock_and_restore_backup(self, client, monkeypatch, capsys):
        client.ready()
        lock = client.workbook.path.with_name(f"~${client.workbook.path.name}")
        lock.write_bytes(b"")
        self.run(monkeypatch, "clear-lock")
        assert not lock.exists()

        client.workbook.path.write_bytes(b"damaged")
        self.run(monkeypatch, "restore-backup")

        assert "Restored" in capsys.readouterr().out
        assert client.get_meta() == {}

    def test_dedupe_round_trip_through_files(self, client, tmp_path, monkeypatch, capsys):
        client.append_jobs([make_job(1)])
        raw_file, out_file = tmp_path / "raw.json", tmp_path / "deduped.json"
        raw_file.write_text(
            json.dumps({"jobs": [make_job(1), make_job(2, title="Z")]}), encoding="utf-8"
        )

        self.run(monkeypatch, "dedupe", "--in", str(raw_file), "--out", str(out_file))

        result = json.loads(out_file.read_text(encoding="utf-8"))
        assert [j["job_id"] for j in result["new"]] == ["2"]
        assert "1 new, 0 reposts skipped" in capsys.readouterr().out


OLD_SHEET = "https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789/edit"


@pytest.fixture
def old_profile(profile_dir):
    """A profile from the Google Sheets version: legacy settings, no workbook yet."""
    (profile_dir / "config.yaml").write_text(
        f'locations: ["Berlin"]\nsheet: "{OLD_SHEET}"\nauth: service_account\n', encoding="utf-8"
    )
    return profile_dir


class TestProfilesFromTheGoogleVersion:
    def test_init_refuses_instead_of_starting_an_empty_list(self, old_profile):
        client = sheet.SheetClient("jane")

        with pytest.raises(sheet.OldGoogleListError) as error:
            client.init()

        message = str(error.value)
        assert OLD_SHEET in message
        assert "File > Download > Microsoft Excel (.xlsx)" in message
        assert str(client.workbook.path) in message
        assert "BEFORE" in message
        assert "init --fresh" in message
        assert not client.workbook.path.exists()
        assert config.read_state("jane") == {}

    def test_every_other_command_says_the_same(self, old_profile):
        with pytest.raises(WorkbookMissing, match="Google sheet before"):
            sheet.SheetClient("jane").get_meta()

    def test_fresh_starts_an_empty_list_on_purpose(self, old_profile):
        client = sheet.SheetClient("jane")
        client.init(fresh=True)
        assert client.workbook.exists()
        assert config.read_state("jane")["workbook"] == str(client.workbook.path)

    def test_the_exported_sheet_is_picked_up_by_init_and_its_history_kept(self, old_profile):
        client = sheet.SheetClient("jane")
        export_like(client.workbook.path, feedback={"7": "Not for me||too junior"})

        client.init()  # no --fresh needed: the list is where it should be

        assert [r["Job ID"] for r in client.workbook.get_records(JOBS_TAB)] == ["7"]

    def test_a_profile_known_only_by_its_old_state_file_is_protected_too(self, profile_dir):
        config.write_state("jane", {"sheet_id": "1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789"})
        with pytest.raises(sheet.OldGoogleListError, match="docs.google.com/spreadsheets/d/1AbC"):
            sheet.SheetClient("jane").init()

    def test_without_any_address_the_message_points_at_google_drive(self, profile_dir):
        (profile_dir / "config.yaml").write_text(
            'locations: ["Berlin"]\nauth: oauth\n', encoding="utf-8"
        )
        with pytest.raises(sheet.OldGoogleListError, match="Google Drive"):
            sheet.SheetClient("jane").init()

    def test_a_new_profile_is_not_affected(self, profile_dir):
        sheet.SheetClient("jane").init()

    def test_a_missing_file_with_a_recorded_workbook_keeps_its_own_explanation(self, old_profile):
        config.write_state("jane", {"workbook": "somewhere.xlsx"})
        with pytest.raises(WorkbookMissing, match="agent-state.json"):
            sheet.SheetClient("jane").init()


def export_like(path, feedback=None):
    """What File > Download > Microsoft Excel gives for a sheet the Google version made: every
    tab visible, the learning step's memory as one JSON blob in AI_MetaData, no AI_StatusSync."""
    book = openpyxl.Workbook()
    jobs = book.active
    jobs.title = JOBS_TAB
    jobs.append(JOBS_HEADERS)
    jobs.append(["2026-10-01", 80, "7", "Analyst", "Acme", "Berlin", "2026-09-30", "", "ok"])
    book.create_sheet(RULES_TAB).append(TAB_HEADERS[RULES_TAB])
    meta = book.create_sheet(META_TAB)
    meta.append(["Key", "Value"])
    meta.append(["last_run_iso", "2026-10-01T00:00:00+00:00"])
    if feedback is not None:
        meta.append([LEGACY_STATUS_SNAPSHOT_KEY, json.dumps(feedback)])
    book.save(path)


class TestImportedList:
    def test_old_feedback_is_not_learned_again(self, profile_dir):
        client = sheet.SheetClient("jane")
        export_like(client.workbook.path, feedback={"7": "Not for me||too junior"})
        edit_cells(client, JOBS_TAB, 2, Status="Not for me", Reason="too junior")

        client.init()

        assert client.workbook.get_records(STATUS_SYNC_TAB) == [
            {"Job ID": "7", "Status": "Not for me", "Reason": "too junior"}
        ]
        assert client.get_changed_status_rows() == []

    def test_without_the_migration_that_feedback_would_be_new(self, profile_dir):
        client = sheet.SheetClient("jane")
        export_like(client.workbook.path)  # no snapshot in the metadata
        edit_cells(client, JOBS_TAB, 2, Status="Not for me", Reason="too junior")
        client.init()
        assert [c["job_id"] for c in client.get_changed_status_rows()] == ["7"]

    def test_the_migration_runs_once_and_never_overwrites_newer_feedback(self, profile_dir):
        client = sheet.SheetClient("jane")
        export_like(client.workbook.path, feedback={"7": "Not for me||too junior"})
        client.init()
        client.workbook.upsert_by_key(
            STATUS_SYNC_TAB,
            "Job ID",
            [{"Job ID": "7", "Status": "Applied", "Reason": ""}],
            append_missing=False,
        )

        client.init()

        assert client.workbook.get_records(STATUS_SYNC_TAB)[0]["Status"] == "Applied"

    @pytest.mark.parametrize("blob", ["not json", "[1, 2]", "{}", '{"7": "no separator"}'])
    def test_an_unusable_snapshot_is_ignored(self, profile_dir, blob):
        client = sheet.SheetClient("jane")
        export_like(client.workbook.path, feedback={})
        edit_cells(client, META_TAB, 3, Value=blob)
        client.init()
        assert client.workbook.get_records(STATUS_SYNC_TAB) == []

    def test_the_bookkeeping_tabs_are_hidden_but_rules_stay_visible(self, profile_dir):
        client = sheet.SheetClient("jane")
        export_like(client.workbook.path, feedback={})

        client.init()

        states = {ws.title: ws.sheet_state for ws in openpyxl.load_workbook(client.workbook.path)}
        assert states[META_TAB] == states[STATUS_SYNC_TAB] == "hidden"
        assert states[RULES_TAB] == states[JOBS_TAB] == "visible"
        assert states[PREFERENCES_TAB] == "visible"


class TestInitWhenNothingIsNeeded:
    def test_it_does_not_save_or_complain_when_the_list_is_already_complete(self, client):
        client.workbook.path.with_name(f"~${client.workbook.path.name}").write_bytes(b"")
        before = client.workbook.path.read_bytes()

        client.init()  # open in Excel, but nothing to add

        assert client.workbook.path.read_bytes() == before

    def test_a_dropdown_the_person_added_elsewhere_survives(self, client):
        book = openpyxl.load_workbook(client.workbook.path)
        ws = book[JOBS_TAB]
        mine = DataValidation(type="list", formula1='"x,y"', allow_blank=True)
        mine.add("N2:N500")
        ws.add_data_validation(mine)
        book.save(client.workbook.path)

        client.init()
        client.apply_status_validation()

        found = {
            str(v.sqref): v.formula1
            for v in openpyxl.load_workbook(client.workbook.path)[
                JOBS_TAB
            ].data_validations.dataValidation
        }
        assert found["N2:N500"] == '"x,y"'
        assert len(found) == 2


class TestInputFilesAndMessages:
    def run(self, monkeypatch, *argv):
        monkeypatch.setattr("sys.argv", ["sheet.py", "--profile", "jane", *argv])
        run_cli(sheet.main)

    def test_add_rules_reports_how_many_were_really_added(
        self, client, tmp_path, monkeypatch, capsys
    ):
        rules = tmp_path / "rules.json"
        rules.write_text(json.dumps([{"rule": "A", "type": "Block"}] * 2), encoding="utf-8")

        self.run(monkeypatch, "add-rules", "--in", str(rules))
        self.run(monkeypatch, "add-rules", "--in", str(rules))

        out = capsys.readouterr().out.splitlines()
        assert out == [
            "Added 1 rules to AI_Rules (1 already there)",
            "Added 0 rules to AI_Rules (2 already there)",
        ]

    @pytest.mark.parametrize(
        "content", ['[{"type": "Block"}]', '[{"rule": "A"}]', '["just text"]', '{"rule": "A"}']
    )
    def test_malformed_rules_are_a_plain_error(self, client, tmp_path, monkeypatch, content):
        rules = tmp_path / "rules.json"
        rules.write_text(content, encoding="utf-8")
        before = client.workbook.path.read_bytes()

        with pytest.raises(SystemExit) as error:
            self.run(monkeypatch, "add-rules", "--in", str(rules))

        assert "Traceback" not in str(error.value)
        assert "rule" in str(error.value).lower()
        assert client.workbook.path.read_bytes() == before

    def test_update_meta_without_an_equals_sign_is_a_plain_error(self, client, monkeypatch):
        with pytest.raises(SystemExit, match="KEY=VALUE.*foo"):
            self.run(monkeypatch, "update-meta", "--set", "foo")

    def test_a_missing_input_file_is_a_plain_error(self, client, tmp_path, monkeypatch):
        with pytest.raises(SystemExit, match="Couldn't read .*nope.json"):
            self.run(monkeypatch, "append-jobs", "--in", str(tmp_path / "nope.json"))

    def test_an_input_file_that_is_not_json_is_a_plain_error(self, client, tmp_path, monkeypatch):
        broken = tmp_path / "broken.json"
        broken.write_text("{not json", encoding="utf-8")
        with pytest.raises(SystemExit, match="isn't valid JSON"):
            self.run(monkeypatch, "append-jobs", "--in", str(broken))

    def test_a_job_without_an_id_is_a_plain_error_and_nothing_is_written(
        self, client, tmp_path, monkeypatch
    ):
        jobs = tmp_path / "jobs.json"
        jobs.write_text(json.dumps([make_job(1), {"title": "no id"}]), encoding="utf-8")
        before = client.workbook.path.read_bytes()

        with pytest.raises(SystemExit, match="no job_id"):
            self.run(monkeypatch, "append-jobs", "--in", str(jobs))

        assert client.workbook.path.read_bytes() == before

    def test_mark_status_synced_reads_the_changed_status_file(
        self, client, tmp_path, monkeypatch, capsys
    ):
        client.append_jobs([make_job(1)])
        edit_cells(client, JOBS_TAB, 2, Status="Applied", Reason="sent CV")
        changed = tmp_path / "changed.json"

        self.run(monkeypatch, "changed-status", "--out", str(changed))
        self.run(monkeypatch, "mark-status-synced", "--in", str(changed))

        assert "Synced status for 1 jobs" in capsys.readouterr().out
        assert client.get_changed_status_rows() == []

    def test_restore_backup_says_where_the_replaced_file_went(self, client, monkeypatch, capsys):
        client.ready()
        client.workbook.path.write_bytes(b"damaged")

        self.run(monkeypatch, "restore-backup")

        out = capsys.readouterr().out
        assert "Restored" in out
        assert "before-restore" in out

    def test_a_backup_that_cannot_be_made_stops_ready_with_a_plain_message(self, client):
        client.workbook.backups_dir.write_text("in the way", encoding="utf-8")
        with pytest.raises(WorkbookError, match="Couldn't make a backup"):
            client.ready()

    def test_messages_carry_the_real_profile_and_command(self, client):
        client.workbook.path.write_bytes(b"junk")
        with pytest.raises(WorkbookCorrupt) as error:
            client.ready()
        assert "uv run python scripts/sheet.py --profile jane restore-backup" in str(error.value)

    def test_the_missing_file_message_has_the_full_command(self, client):
        client.workbook.path.unlink()
        with pytest.raises(WorkbookMissing) as error:
            client.get_meta()
        assert "`uv run python scripts/sheet.py --profile jane restore-backup`" in str(error.value)
