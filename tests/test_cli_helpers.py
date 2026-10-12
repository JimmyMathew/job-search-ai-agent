import io
import os
import stat
import sys
from pathlib import Path

import config
import doctor
import new_profile
import openpyxl
import pipeline
import pytest
import sheet
import workbook as workbook_module
from config import UserFacingError, looks_like_resume_file
from util import run_cli
from workbook import WorkbookError

OLD_SHEET_URL = "https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789/edit"


class TestRunCli:
    def test_user_facing_errors_end_with_the_plain_message(self):
        def main():
            raise UserFacingError("Please close the file.")

        with pytest.raises(SystemExit) as exit_info:
            run_cli(main)
        assert str(exit_info.value) == "Please close the file."

    def test_an_unexpected_error_becomes_one_plain_sentence(self, monkeypatch):
        monkeypatch.delenv("JOB_AGENT_DEBUG", raising=False)

        def main():
            raise ValueError("a real bug")

        with pytest.raises(SystemExit) as exit_info:
            run_cli(main)

        message = str(exit_info.value)
        assert "JOB_AGENT_DEBUG=1" in message
        assert "a real bug" not in message
        assert exit_info.value.code != 0

    def test_debug_mode_shows_the_real_error(self, monkeypatch):
        monkeypatch.setenv("JOB_AGENT_DEBUG", "1")

        def main():
            raise ValueError("a real bug")

        with pytest.raises(ValueError, match="a real bug"):
            run_cli(main)

    def test_output_is_utf8_even_when_the_console_uses_a_legacy_code_page(self, monkeypatch):
        raw = io.BytesIO()
        console = io.TextIOWrapper(raw, encoding="cp1252", errors="strict")
        monkeypatch.setattr(sys, "stdout", console)

        run_cli(lambda: print("مرحبا 😀"))

        console.flush()
        assert raw.getvalue().decode("utf-8").strip() == "مرحبا 😀"

    def test_a_stream_that_cannot_be_reconfigured_is_left_alone(self, monkeypatch):
        class Plain:
            def write(self, text):
                return len(text)

        monkeypatch.setattr(sys, "stdout", Plain())
        run_cli(lambda: None)

    def test_known_errors_are_user_facing(self):
        assert issubclass(config.ProfileError, UserFacingError)
        assert issubclass(WorkbookError, UserFacingError)


class TestLooksLikeResumeFile:
    @pytest.mark.parametrize("name", ["Jane Resume.pdf", "cv.DOCX", "notes.txt"])
    def test_documents_are_accepted(self, tmp_path, name):
        assert looks_like_resume_file(tmp_path / name)

    @pytest.mark.parametrize(
        "name", ["profile.md", "README.md", "readme.txt", "config.yaml", "x.png"]
    )
    def test_our_own_files_and_other_types_are_not(self, tmp_path, name):
        assert not looks_like_resume_file(tmp_path / name)


class TestNewProfileResume:
    def test_rejects_a_file_that_would_clash_with_ours(self, tmp_path, monkeypatch, capsys):
        monkeypatch.setattr(config, "PROFILES_DIR", tmp_path / "profiles")
        monkeypatch.setattr(new_profile, "PROFILES_DIR", tmp_path / "profiles")
        bad = tmp_path / "profile.md"
        bad.write_text("x", encoding="utf-8")
        monkeypatch.setattr("sys.argv", ["new_profile.py", "jane", "--resume", str(bad)])
        with pytest.raises(SystemExit, match="doesn't look like a resume"):
            new_profile.main()
        assert not (tmp_path / "profiles" / "jane").exists()

    def test_drop_folder_ignores_the_readme_and_needs_exactly_one_file(self, tmp_path, monkeypatch):
        drop = tmp_path / "resume"
        drop.mkdir()
        monkeypatch.setattr(new_profile, "DROP_FOLDER", drop)
        (drop / "README.txt").write_text("hi", encoding="utf-8")
        assert new_profile.drop_folder_resume() is None
        (drop / "Jane.pdf").write_text("x", encoding="utf-8")
        assert new_profile.drop_folder_resume().name == "Jane.pdf"
        (drop / "Other.docx").write_text("x", encoding="utf-8")
        assert new_profile.drop_folder_resume() is None


class TestSplitScorable:
    def test_partitions_in_one_pass_and_keeps_order(self):
        jobs = [
            {"job_id": "1", "description": "text"},
            {"job_id": "2", "description": ""},
            {"job_id": "3", "description": "text", "description_unavailable": True},
            {"job_id": "4", "description": "more"},
        ]
        scorable, unavailable = pipeline.split_scorable(jobs)
        assert [j["job_id"] for j in scorable] == ["1", "4"]
        assert [j["job_id"] for j in unavailable] == ["2", "3"]


class TestSyncedFolders:
    @pytest.mark.parametrize(
        "path",
        [
            "C:/Users/Jane/OneDrive/Jobs/list.xlsx",
            "C:/Users/Jane/OneDrive - Acme Ltd/list.xlsx",
            "/home/jane/Dropbox/list.xlsx",
            "G:/My Drive/Google Drive/list.xlsx",
        ],
    )
    def test_sync_folders_are_recognised(self, path):
        assert doctor.synced_folder_name(Path(path)) is not None

    def test_ordinary_folders_are_not(self):
        assert doctor.synced_folder_name(Path("C:/Users/Jane/Documents/list.xlsx")) is None


class TestDoctorWorkbookCheck:
    @pytest.fixture
    def reports(self, monkeypatch):
        collected = []
        monkeypatch.setattr(
            doctor, "report", lambda level, msg, fix="": collected.append((level, msg, fix))
        )
        return collected

    @pytest.fixture
    def profile(self, tmp_path, monkeypatch):
        monkeypatch.setattr(config, "PROFILES_DIR", tmp_path)
        (tmp_path / "jane").mkdir()
        return "jane"

    def config(self, **overrides):
        return config.normalize_config({"locations": ["Berlin"], **overrides}, "jane")

    def create(self, profile, cfg):
        """The job list as `sheet.py init` makes it, for whatever path the config names."""
        workbook = workbook_module.Workbook(
            config.workbook_path(profile, cfg), config.backups_dir(profile)
        )
        workbook.ensure_tabs(sheet.TAB_SPECS)
        return workbook

    def levels(self, reports):
        return [level for level, _, _ in reports]

    def test_a_missing_list_says_how_to_create_it(self, profile, reports):
        doctor.check_workbook(profile, self.config())
        ((level, message, fix),) = reports
        assert level == "FIX"
        assert "Job Search - jane.xlsx" in message
        assert "sheet.py --profile jane init" in fix
        assert "Google" not in fix

    def test_an_old_google_profile_is_told_how_to_bring_its_data_over(self, profile, reports):
        cfg = self.config(sheet=OLD_SHEET_URL, auth="oauth")
        doctor.check_workbook(profile, cfg)

        fix, warn = reports
        assert fix[0] == "FIX"
        assert OLD_SHEET_URL in fix[1]
        assert "File > Download > Microsoft Excel (.xlsx)" in fix[2]
        assert str(config.workbook_path(profile, cfg)) in fix[2]
        assert "BEFORE init" in fix[2]
        assert "init --fresh" in fix[2]
        assert warn[0] == "WARN"
        assert "`sheet`, `auth`" in warn[1]
        assert "no longer used" in warn[1]

    def test_old_settings_are_not_called_deletable_until_the_list_exists(self, profile, reports):
        cfg = self.config(sheet=OLD_SHEET_URL)
        doctor.check_workbook(profile, cfg)
        assert "delete" not in reports[-1][2]

        reports.clear()
        self.create(profile, cfg)
        doctor.check_workbook(profile, cfg)
        assert "You can delete those lines" in reports[-1][2]

    def test_an_old_sheet_known_only_from_the_saved_state_is_found_too(self, profile, reports):
        config.write_state(profile, {"sheet_id": "1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789"})
        doctor.check_workbook(profile, self.config())
        assert "docs.google.com/spreadsheets/d/1AbCdEf" in reports[0][1]

    def test_an_old_profile_without_a_known_address_is_pointed_to_google_drive(
        self, profile, reports
    ):
        doctor.check_workbook(profile, self.config(auth="oauth"))
        assert "Google Drive" in reports[0][1]

    def test_a_healthy_list_passes_and_shows_where_it_is(self, profile, reports):
        cfg = self.config()
        workbook = self.create(profile, cfg)

        doctor.check_workbook(profile, cfg)

        assert reports == [("OK", f"Job list: {workbook.path}", "")]

    def test_missing_tabs_are_a_fix_that_only_adds(self, profile, reports):
        cfg = self.config()
        workbook = self.create(profile, cfg)
        book = openpyxl.load_workbook(workbook.path)
        del book["AI_Rules"]
        book.save(workbook.path)

        doctor.check_workbook(profile, cfg)

        ((level, message, fix),) = reports
        assert level == "FIX"
        assert "AI_Rules" in message
        assert "init" in fix

    def test_a_damaged_file_is_a_fix_with_the_restore_command_not_a_traceback(
        self, profile, reports
    ):
        cfg = self.config()
        workbook = self.create(profile, cfg)
        workbook.path.write_bytes(b"not an excel file")

        doctor.check_workbook(profile, cfg)

        ((level, message, fix),) = reports
        assert level == "FIX"
        assert "can't be opened" in message
        assert "restore-backup" in fix

    def test_a_list_open_in_excel_is_a_warning(self, profile, reports):
        cfg = self.config()
        workbook = self.create(profile, cfg)
        workbook.path.with_name(f"~${workbook.path.name}").write_bytes(b"")

        doctor.check_workbook(profile, cfg)

        assert self.levels(reports) == ["OK", "WARN"]
        assert "open in Excel" in reports[1][1]

    def test_a_list_in_a_synced_folder_is_a_warning(self, profile, reports, tmp_path):
        target = tmp_path / "OneDrive - Acme" / "jobs.xlsx"
        cfg = self.config(workbook=str(target))
        self.create(profile, cfg)

        doctor.check_workbook(profile, cfg)

        assert self.levels(reports) == ["OK", "WARN"]
        assert "OneDrive - Acme" in reports[1][1]

    def test_copies_that_look_like_conflicts_are_listed(self, profile, reports):
        cfg = self.config()
        workbook = self.create(profile, cfg)
        for name in ("Job Search - jane (1).xlsx", "Job Search - jane-LAPTOP.xlsx", "Other.xlsx"):
            (workbook.path.parent / name).write_bytes(b"x")

        doctor.check_workbook(profile, cfg)

        assert self.levels(reports) == ["OK", "WARN"]
        assert "(1).xlsx" in reports[1][2]
        assert "LAPTOP" in reports[1][2]
        assert "Other.xlsx" not in reports[1][2]

    def test_find_conflict_copies_ignores_the_list_itself_and_lock_files(self, tmp_path):
        path = tmp_path / "Job Search - jane.xlsx"
        for name in (path.name, "~$Job Search - jane.xlsx", ".Job Search - jane.saving.xlsx"):
            (tmp_path / name).write_bytes(b"x")
        assert doctor.find_conflict_copies(path) == []

    def test_find_conflict_copies_in_a_folder_that_does_not_exist(self, tmp_path):
        assert doctor.find_conflict_copies(tmp_path / "nope" / "a.xlsx") == []


class TestDoctorWorkbookExtras:
    @pytest.fixture
    def reports(self, monkeypatch):
        collected = []
        monkeypatch.setattr(
            doctor, "report", lambda level, msg, fix="": collected.append((level, msg, fix))
        )
        return collected

    @pytest.fixture
    def listed(self, tmp_path, monkeypatch):
        monkeypatch.setattr(config, "PROFILES_DIR", tmp_path)
        (tmp_path / "jane").mkdir()
        cfg = config.normalize_config({"locations": ["Berlin"]}, "jane")
        workbook = workbook_module.Workbook(
            config.workbook_path("jane", cfg), config.backups_dir("jane")
        )
        workbook.ensure_tabs(sheet.TAB_SPECS)
        return cfg, workbook

    def test_a_read_only_list_is_a_fix_with_the_way_out(self, listed, reports):
        cfg, workbook = listed
        workbook.path.chmod(stat.S_IREAD)
        try:
            doctor.check_workbook("jane", cfg)
        finally:
            workbook.path.chmod(stat.S_IWRITE | stat.S_IREAD)

        level, message, fix = reports[1]
        assert level == "FIX"
        assert "read-only" in message
        assert "Properties" in fix

    def test_leftover_temp_files_are_a_warning_that_names_the_cleanup(self, listed, reports):
        cfg, workbook = listed
        leftover = workbook.path.with_name(f".{workbook.path.stem}.abc123.saving.xlsx")
        leftover.write_bytes(b"half")

        doctor.check_workbook("jane", cfg)

        level, message, fix = reports[-1]
        assert level == "WARN"
        assert leftover.name in fix
        assert "clear-lock" in fix

    def test_an_update_in_progress_is_a_warning_not_a_failure(self, listed, reports, monkeypatch):
        cfg, workbook = listed
        monkeypatch.setattr(workbook_module, "LOCK_WAIT_SECONDS", 0.2)
        workbook.path.with_name(f".{workbook.path.stem}.agent-lock").write_text(
            f"{os.getpid()} 1", encoding="utf-8"
        )

        doctor.check_workbook("jane", cfg)

        assert [level for level, _, _ in reports] == ["WARN"]

    @pytest.mark.parametrize(
        "name",
        [
            "Job Search - jane (1).xlsx",
            "Job Search - jane (12).xlsx",
            "Job Search - jane - Copy.xlsx",
            "Job Search - jane - Copy (2).xlsx",
            "Job Search - jane-DESKTOP-ABC123.xlsx",
            "Job Search - jane (Jimmy's conflicted copy 2026-10-12).xlsx",
        ],
    )
    def test_names_sync_programs_give_copies_are_recognised(self, tmp_path, name):
        path = tmp_path / "Job Search - jane.xlsx"
        (tmp_path / name).write_bytes(b"x")
        assert [p.name for p in doctor.find_conflict_copies(path)] == [name]

    @pytest.mark.parametrize(
        "name",
        [
            "Job Search - jane-doe.xlsx",  # another profile whose name starts the same way
            "Job Search - jane-2.xlsx",
            "Job Search - janet.xlsx",
            "Job Search - jane.xls",
            "Job Search - jane (draft).xlsx",
            "Notes.xlsx",
        ],
    )
    def test_other_files_that_merely_start_alike_are_not_copies(self, tmp_path, name):
        path = tmp_path / "Job Search - jane.xlsx"
        (tmp_path / name).write_bytes(b"x")
        assert doctor.find_conflict_copies(path) == []
