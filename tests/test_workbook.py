import os
import stat
import subprocess
import sys
import threading
import time
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import pytest
import workbook as workbook_module
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table
from workbook import (
    LOCKED_PREFIX,
    MAX_BACKUPS,
    TabSpec,
    Workbook,
    WorkbookBusy,
    WorkbookCorrupt,
    WorkbookError,
    WorkbookLocked,
    WorkbookMissing,
    WorkbookReadOnly,
    is_linkable,
    normalise,
    storable,
)

JOB_HEADERS = ["Job ID", "Score", "Link", "Status", "Notes"]
SPECS = [
    TabSpec("Jobs", JOB_HEADERS, widths={"Score": 8}),
    TabSpec("Prefs", ["Criterion", "Value"], seed_rows=[("Salary?", ""), ("Remote?", "")]),
    TabSpec("Rules", ["Rule", "Type"]),
    TabSpec("Hidden", ["Key", "Value"], hidden=True),
]


@pytest.fixture
def wb(tmp_path):
    return Workbook(tmp_path / "Job Search - jane.xlsx", tmp_path / "backups")


@pytest.fixture
def created(wb):
    wb.ensure_tabs(SPECS)
    return wb


def raw(wb):
    """The saved file read straight with openpyxl, independent of the class under test."""
    return openpyxl.load_workbook(wb.path)


def job(job_id, **extra):
    return {"Job ID": job_id, "Score": "", "Link": "", "Status": "New", "Notes": "", **extra}


def snapshot(path):
    return path.read_bytes(), path.stat().st_mtime_ns


class TestCreate:
    def test_tabs_exist_in_order_with_hidden_ones_hidden(self, created):
        book = raw(created)
        assert book.sheetnames == ["Jobs", "Prefs", "Rules", "Hidden"]
        assert [ws.sheet_state for ws in book] == ["visible", "visible", "visible", "hidden"]

    def test_jobs_is_the_active_tab(self, created):
        book = raw(created)
        assert book.active.title == "Jobs"
        assert [ws.title for ws in book if ws.sheet_view.tabSelected] == ["Jobs"]

    def test_headers_are_styled_and_the_top_row_frozen(self, created):
        ws = raw(created)["Jobs"]
        assert [c.value for c in ws[1]] == JOB_HEADERS
        assert ws.freeze_panes == "A2"
        assert ws["A1"].font.bold
        assert ws.column_dimensions["B"].width == 8

    def test_seed_rows_only_go_into_a_new_tab(self, created):
        assert created.get_values("Prefs") == [
            ["Criterion", "Value"],
            ["Salary?", ""],
            ["Remote?", ""],
        ]

    def test_autofilter_covers_the_data_after_every_save(self, created):
        created.append_rows("Jobs", [job("1"), job("2")], unique_header="Job ID")
        assert raw(created)["Jobs"].auto_filter.ref == "A1:E3"

    def test_a_missing_file_is_not_created_by_reading(self, wb):
        with pytest.raises(WorkbookMissing):
            wb.get_values("Jobs")
        assert not wb.path.exists()

    def test_no_leftover_temp_file_after_a_save(self, created):
        assert sorted(p.name for p in created.path.parent.iterdir()) == [created.path.name]


class TestEnsureTabsIsAdditive:
    def test_running_twice_changes_nothing_and_keeps_what_the_person_typed(self, created):
        book = raw(created)
        book["Prefs"]["B2"] = "60,000 a year"
        book.save(created.path)
        before = snapshot(created.path)

        created.ensure_tabs(SPECS)

        assert snapshot(created.path) == before
        assert created.get_records("Prefs")[0]["Value"] == "60,000 a year"

    def test_a_missing_tab_is_added_without_touching_the_others(self, created):
        book = raw(created)
        del book["Rules"]
        book.save(created.path)
        created.append_rows("Jobs", [job("1")], unique_header="Job ID")

        created.ensure_tabs(SPECS)

        assert "Rules" in raw(created).sheetnames
        assert created.get_records("Jobs")[0]["Job ID"] == "1"

    def test_a_missing_header_is_added_at_the_end(self, created):
        created.append_rows("Jobs", [job("1", Notes="keep")], unique_header="Job ID")
        wider = [TabSpec("Jobs", [*JOB_HEADERS, "Salary"])]

        created.ensure_tabs(wider)

        assert created.get_headers("Jobs") == [*JOB_HEADERS, "Salary"]
        assert created.get_records("Jobs")[0]["Notes"] == "keep"

    def test_a_reordered_sheet_keeps_its_order_and_only_gains_what_is_missing(self, created):
        book = raw(created)
        ws = book["Jobs"]
        ws["A1"], ws["B1"] = "Score", "Job ID"
        book.save(created.path)

        created.ensure_tabs(SPECS)

        assert created.get_headers("Jobs") == ["Score", "Job ID", *JOB_HEADERS[2:]]


class TestAppendRows:
    def test_cells_are_placed_by_header_name_not_position(self, created):
        created.append_rows("Jobs", [{"Notes": "n", "Job ID": "7"}])
        (record,) = created.get_records("Jobs")
        assert (record["Job ID"], record["Notes"], record["Status"]) == ("7", "n", "")

    def test_unknown_headers_are_skipped(self, created):
        assert created.append_rows("Jobs", [{"Job ID": "7", "Nope": "x"}]) == 1
        assert created.get_headers("Jobs") == JOB_HEADERS

    def test_rows_with_a_known_key_are_skipped_so_a_retry_adds_nothing(self, created):
        rows = [job("1"), job("2"), job("2")]
        assert created.append_rows("Jobs", rows, unique_header="Job ID") == 2
        before = snapshot(created.path)
        assert created.append_rows("Jobs", rows, unique_header="Job ID") == 0
        assert snapshot(created.path) == before

    def test_new_rows_go_after_the_last_row_with_a_key_not_after_formatted_rows(self, created):
        book = raw(created)
        book["Jobs"]["A40"].fill = openpyxl.styles.PatternFill("solid", fgColor="FFFF00")
        book.save(created.path)

        created.append_rows("Jobs", [job("1")], unique_header="Job ID")

        assert raw(created)["Jobs"]["A2"].value == "1"

    def test_a_row_without_its_key_is_rejected(self, created):
        with pytest.raises(WorkbookError, match="Job ID"):
            created.append_rows("Jobs", [{"Notes": "x"}], unique_header="Job ID")

    def test_a_deleted_key_column_is_explained(self, created):
        book = raw(created)
        book["Jobs"]["A1"] = None
        book.save(created.path)
        with pytest.raises(WorkbookError, match="no 'Job ID' column"):
            created.append_rows("Jobs", [job("1")], unique_header="Job ID")

    def test_a_deleted_tab_is_explained(self, created):
        book = raw(created)
        del book["Rules"]
        book.save(created.path)
        with pytest.raises(WorkbookError, match="no 'Rules' tab"):
            created.get_values("Rules")


class TestUpsertByKey:
    def test_only_the_given_columns_change(self, created):
        created.append_rows("Jobs", [job("1", Notes="mine", Status="Applied")], "Job ID")

        written = created.upsert_by_key("Jobs", "Job ID", [{"Job ID": "1", "Score": 80}], False)

        assert written == 1
        record = created.get_records("Jobs")[0]
        assert (record["Score"], record["Notes"], record["Status"]) == ("80", "mine", "Applied")

    def test_an_unknown_key_is_never_appended_when_not_asked(self, created):
        created.append_rows("Jobs", [job("1")], "Job ID")
        before = snapshot(created.path)

        assert created.upsert_by_key("Jobs", "Job ID", [{"Job ID": "9", "Score": 5}], False) == 0

        assert snapshot(created.path) == before

    def test_an_unknown_key_is_appended_when_asked_and_only_once(self, created):
        rows = [{"Job ID": "9", "Score": 5}, {"Job ID": "9", "Score": 6}]
        assert created.upsert_by_key("Jobs", "Job ID", rows, True) == 2
        assert [r["Score"] for r in created.get_records("Jobs")] == ["6"]

    def test_a_sorted_sheet_still_updates_the_right_rows(self, created):
        created.append_rows("Jobs", [job(str(i)) for i in range(1, 6)], "Job ID")
        book = raw(created)
        ws = book["Jobs"]
        ids = [ws.cell(r, 1).value for r in range(2, 7)]
        for row, job_id in zip(range(2, 7), reversed(ids), strict=True):
            ws.cell(row, 1).value = job_id  # the person sorted Job ID descending
        book.save(created.path)

        created.upsert_by_key("Jobs", "Job ID", [{"Job ID": "2", "Notes": "second"}], False)

        by_id = {r["Job ID"]: r["Notes"] for r in created.get_records("Jobs")}
        assert by_id == {"1": "", "2": "second", "3": "", "4": "", "5": ""}

    def test_writing_a_blank_clears_the_cell(self, created):
        created.append_rows("Jobs", [job("1", Score=80)], "Job ID")
        created.upsert_by_key("Jobs", "Job ID", [{"Job ID": "1", "Score": ""}], False)
        assert created.get_records("Jobs")[0]["Score"] == ""


class TestSession:
    def test_several_operations_share_one_save(self, created, monkeypatch):
        saves = []
        real_replace = os.replace
        monkeypatch.setattr(
            workbook_module.os, "replace", lambda *a: (saves.append(a), real_replace(*a))
        )

        with created.session():
            created.append_rows("Jobs", [job("1")], "Job ID")
            created.append_rows("Jobs", [job("2")], "Job ID")
            created.upsert_by_key("Jobs", "Job ID", [{"Job ID": "1", "Score": 5}], False)
            assert created.get_records("Jobs")[0]["Score"] == "5"  # visible before it is saved

        assert len(saves) == 1
        assert len(created.get_records("Jobs")) == 2

    def test_an_error_inside_a_session_saves_nothing(self, created):
        before = snapshot(created.path)
        with pytest.raises(RuntimeError), created.session():
            created.append_rows("Jobs", [job("1")], "Job ID")
            raise RuntimeError("boom")
        assert snapshot(created.path) == before

    def test_reading_never_saves(self, created):
        before = snapshot(created.path)
        created.get_values("Jobs")
        assert snapshot(created.path) == before


class TestCellSafety:
    @pytest.mark.parametrize(
        "text", ["=1+1", '=HYPERLINK("http://evil","x")', "=cmd|' /C calc'!A0"]
    )
    def test_text_starting_with_equals_is_stored_as_text_not_a_formula(self, created, text):
        created.append_rows("Jobs", [job("1", Notes=text)], "Job ID")

        cell = raw(created)["Jobs"]["E2"]
        assert cell.value == text
        assert cell.data_type == "s"
        assert created.get_records("Jobs")[0]["Notes"] == text

    def test_illegal_control_characters_are_removed(self, created):
        created.append_rows("Jobs", [job("1", Notes="bad\x00te\x07xt\n")], "Job ID")
        assert created.get_records("Jobs")[0]["Notes"] == "bad" + "text\n"

    def test_score_is_a_number_and_job_id_stays_text(self, created):
        created.append_rows("Jobs", [job("0012345678901", Score=80)], "Job ID")

        ws = raw(created)["Jobs"]
        assert ws["B2"].value == 80
        assert isinstance(ws["B2"].value, int)
        assert ws["A2"].value == "0012345678901"
        assert ws["A2"].data_type == "s"

    def test_long_text_is_kept_up_to_the_limit(self, created):
        created.append_rows("Jobs", [job("1", Notes="x" * 10_000)], "Job ID")
        assert len(created.get_records("Jobs")[0]["Notes"]) == 10_000

    def test_text_beyond_what_excel_holds_is_cut(self, created):
        created.append_rows("Jobs", [job("1", Notes="x" * 40_000)], "Job ID")
        assert len(created.get_records("Jobs")[0]["Notes"]) == workbook_module.MAX_CELL_CHARS

    def test_dates_stay_iso_text(self, created):
        created.append_rows("Jobs", [job("1", Notes="2026-10-01")], "Job ID")
        cell = raw(created)["Jobs"]["E2"]
        assert cell.value == "2026-10-01"
        assert cell.data_type == "s"


class TestHyperlinks:
    def test_a_valid_link_is_clickable_and_keeps_the_url_as_its_text(self, created):
        url = "https://www.linkedin.com/jobs/view/123"
        created.append_rows("Jobs", [job("1", Link=url)], "Job ID")

        cell = raw(created)["Jobs"]["C2"]
        assert cell.value == url
        assert cell.hyperlink.target == url
        assert cell.style == "Hyperlink"

    @pytest.mark.parametrize("value", ["not a url", "ftp://x/y", "https://x/" + "a" * 2_000, ""])
    def test_anything_else_is_plain_text(self, created, value):
        created.append_rows("Jobs", [job("1", Link=value)], "Job ID")
        assert raw(created)["Jobs"]["C2"].hyperlink is None
        assert created.get_records("Jobs")[0]["Link"] == value

    def test_is_linkable(self):
        assert is_linkable("http://a.b/c?d=1")
        assert not is_linkable("https://a.b/with space")


class TestStatusDropdown:
    def test_exactly_one_list_validation_on_the_column_below_the_header(self, created):
        created.set_list_validation("Jobs", "Status", ["New", "Applied"])
        created.set_list_validation("Jobs", "Status", ["New", "Applied"])

        validations = raw(created)["Jobs"].data_validations.dataValidation
        assert len(validations) == 1
        (validation,) = validations
        assert validation.type == "list"
        assert validation.formula1 == '"New,Applied"'
        assert str(validation.sqref) == "D2:D1048576"
        assert not validation.showDropDown  # True would hide the arrow
        assert validation.allow_blank

    def test_a_dropdown_on_another_column_is_kept(self, created):
        created.set_list_validation("Jobs", "Notes", ["a"])
        created.set_list_validation("Jobs", "Status", ["b"])
        found = {
            str(v.sqref)[0]: v.formula1
            for v in raw(created)["Jobs"].data_validations.dataValidation
        }
        assert found == {"E": '"a"', "D": '"b"'}

    def test_changing_the_values_replaces_only_that_columns_dropdown(self, created):
        created.set_list_validation("Jobs", "Notes", ["a"])
        created.set_list_validation("Jobs", "Status", ["b"])
        created.set_list_validation("Jobs", "Status", ["c", "d"])
        found = {
            str(v.sqref)[0]: v.formula1
            for v in raw(created)["Jobs"].data_validations.dataValidation
        }
        assert found == {"E": '"a"', "D": '"c,d"'}

    def test_asking_for_the_dropdown_it_already_has_changes_nothing(self, created):
        created.set_list_validation("Jobs", "Status", ["a", "b"])
        created.path.with_name(f"~${created.path.name}").write_bytes(b"")  # open: a save would fail
        before = snapshot(created.path)

        created.set_list_validation("Jobs", "Status", ["a", "b"])

        assert snapshot(created.path) == before

    @pytest.mark.parametrize("values", [["a,b"], ["x" * 300]])
    def test_values_excel_cannot_hold_are_rejected(self, created, values):
        with pytest.raises(ValueError, match="Dropdown"):
            created.set_list_validation("Jobs", "Status", values)


class TestReading:
    def test_values_are_normalised_and_blank_rows_skipped(self, created):
        book = raw(created)
        ws = book["Rules"]
        ws.append(["when", date(2026, 10, 1)])
        ws.append([None, None])
        ws.append([80.0, datetime(2026, 10, 2, 13, 30)])
        ws.append([1.5, True])
        book.save(created.path)

        assert created.get_values("Rules") == [
            ["Rule", "Type"],
            ["when", "2026-10-01"],
            ["80", "2026-10-02"],
            ["1.5", "True"],
        ]

    def test_records_are_dicts_by_header(self, created):
        created.append_rows("Rules", [{"Rule": "no sales", "Type": "Block"}])
        assert created.get_records("Rules") == [{"Rule": "no sales", "Type": "Block"}]

    def test_an_empty_tab_has_no_records(self, created):
        assert created.get_records("Hidden") == []

    @pytest.mark.parametrize(
        ("value", "expected"),
        [
            (None, ""),
            (3.0, "3"),
            (3.5, "3.5"),
            (7, "7"),
            ("a", "a"),
            (date(2026, 1, 2), "2026-01-02"),
        ],
    )
    def test_normalise(self, value, expected):
        assert normalise(value) == expected

    def test_storable_keeps_numbers_and_blanks_text(self):
        assert storable(80) == 80
        assert storable("") is None
        assert storable(None) is None
        assert storable(True) == "True"


class TestOpenFileDetection:
    @pytest.mark.parametrize(
        "make_lock_name",
        [
            lambda name: f"~${name}",
            lambda name: f"~${name[2:]}",
            lambda name: f".~lock.{name}#",
        ],
        ids=["excel", "excel-long-name", "libreoffice"],
    )
    def test_a_lock_file_blocks_writes_and_leaves_the_file_untouched(self, created, make_lock_name):
        created.append_rows("Jobs", [job("1")], "Job ID")
        lock = created.path.with_name(make_lock_name(created.path.name))
        lock.write_bytes(b"")
        before = snapshot(created.path)

        assert created.is_open()
        with pytest.raises(WorkbookLocked) as error:
            created.append_rows("Jobs", [job("2")], "Job ID")

        assert str(error.value).startswith(LOCKED_PREFIX)
        assert snapshot(created.path) == before
        assert sorted(p.name for p in created.path.parent.iterdir()) == sorted(
            [created.path.name, lock.name]
        )

    def test_reading_is_allowed_while_it_is_open(self, created):
        created.path.with_name(f"~${created.path.name}").write_bytes(b"")
        assert created.get_headers("Jobs") == JOB_HEADERS

    def test_nothing_to_write_means_no_complaint(self, created):
        created.append_rows("Jobs", [job("1")], "Job ID")
        created.path.with_name(f"~${created.path.name}").write_bytes(b"")
        assert created.append_rows("Jobs", [job("1")], "Job ID") == 0

    def test_an_open_file_is_not_replaced_by_a_new_one(self, wb):
        wb.path.with_name(f"~${wb.path.name}").write_bytes(b"")
        with pytest.raises(WorkbookLocked):
            wb.ensure_tabs(SPECS)
        assert not wb.path.exists()

    def test_closed_means_not_open(self, created):
        assert not created.is_open()
        created.require_closed()

    def test_clear_lock_files_removes_only_lock_files(self, created):
        locks = [
            created.path.with_name(f"~${created.path.name}"),
            created.path.with_name(f".~lock.{created.path.name}#"),
        ]
        for lock in locks:
            lock.write_bytes(b"")
        assert sorted(created.clear_lock_files()) == sorted(locks)
        assert sorted(p.name for p in created.path.parent.iterdir()) == [created.path.name]

    def test_clear_lock_files_with_none_present(self, created):
        assert created.clear_lock_files() == []


class TestSafeSaving:
    def test_a_replace_that_keeps_failing_leaves_the_original_and_no_temp_file(
        self, created, monkeypatch
    ):
        created.append_rows("Jobs", [job("1")], "Job ID")
        before = snapshot(created.path)
        attempts = []

        def refuse(*args):
            attempts.append(args)
            raise PermissionError("held by antivirus")

        monkeypatch.setattr(workbook_module.os, "replace", refuse)
        monkeypatch.setattr(workbook_module.time, "sleep", lambda seconds: None)

        with pytest.raises(WorkbookLocked) as error:
            created.append_rows("Jobs", [job("2")], "Job ID")

        assert str(error.value).startswith(LOCKED_PREFIX)
        assert len(attempts) == workbook_module.REPLACE_RETRIES + 1
        assert snapshot(created.path) == before
        assert [p.name for p in created.path.parent.iterdir() if p.name != "backups"] == [
            created.path.name
        ]

    def test_a_replace_that_succeeds_on_a_retry_saves_normally(self, created, monkeypatch):
        real_replace = os.replace
        calls = []

        def flaky(source, target):
            calls.append(source)
            if len(calls) == 1:
                raise PermissionError("busy")
            real_replace(source, target)

        monkeypatch.setattr(workbook_module.os, "replace", flaky)
        monkeypatch.setattr(workbook_module.time, "sleep", lambda seconds: None)

        created.append_rows("Jobs", [job("1")], "Job ID")

        assert len(calls) == 2
        assert created.get_records("Jobs")[0]["Job ID"] == "1"

    def test_temp_files_are_beside_the_workbook_unique_and_never_named_like_an_excel_lock(
        self, created
    ):
        first, second = created._new_temp_file(), created._new_temp_file()
        try:
            assert first != second
            for temp in (first, second):
                assert not temp.name.startswith("~$")
                assert temp.parent == created.path.parent
                assert temp.name.startswith(f".{created.path.stem}.")
                assert temp.name.endswith(".saving.xlsx")
        finally:
            first.unlink()
            second.unlink()


class TestCorruptFile:
    def test_a_damaged_file_is_reported_and_never_replaced(self, wb):
        wb.path.write_bytes(b"this is not a zip file")
        before = snapshot(wb.path)

        for action in (
            lambda: wb.get_values("Jobs"),
            lambda: wb.ensure_tabs(SPECS),
            lambda: wb.append_rows("Jobs", [job("1")]),
            wb.verify,
        ):
            with pytest.raises(WorkbookCorrupt, match="restore-backup"):
                action()

        assert snapshot(wb.path) == before

    def test_a_zip_that_is_not_a_workbook_is_also_corrupt(self, wb):
        import zipfile

        with zipfile.ZipFile(wb.path, "w") as archive:
            archive.writestr("hello.txt", "hi")
        with pytest.raises(WorkbookCorrupt):
            wb.verify()


class TestBackups:
    def test_a_backup_is_a_copy_named_with_the_time(self, created):
        target = created.backup(now=datetime(2026, 10, 12, 9, 5, 7))
        assert target == created.backups_dir / "Job Search - jane-20261012-090507.xlsx"
        assert target.read_bytes() == created.path.read_bytes()

    def test_only_the_newest_are_kept(self, created):
        start = datetime(2026, 10, 1)
        for i in range(MAX_BACKUPS + 3):
            created.backup(now=start + timedelta(minutes=i))
        names = [p.name for p in created.backup_files()]
        assert len(names) == MAX_BACKUPS
        assert names[0].endswith("20261001-000300.xlsx")
        assert names[-1].endswith("20261001-001200.xlsx")

    def test_other_files_in_the_backup_folder_are_left_alone(self, created):
        created.backups_dir.mkdir()
        stranger = created.backups_dir / "Job Search - jane-notes.xlsx"
        stranger.write_bytes(b"x")
        for i in range(MAX_BACKUPS + 2):
            created.backup(now=datetime(2026, 10, 1) + timedelta(minutes=i))
        assert stranger.exists()
        assert stranger not in created.backup_files()

    def test_restore_puts_the_newest_backup_back(self, created):
        created.append_rows("Jobs", [job("1")], "Job ID")
        created.backup(now=datetime(2026, 10, 1))
        created.append_rows("Jobs", [job("2")], "Job ID")
        created.backup(now=datetime(2026, 10, 2))
        created.path.write_bytes(b"damaged")

        restored, replaced = created.restore_latest_backup(now=datetime(2026, 10, 3, 8, 0, 0))

        assert restored.name.endswith("20261002-000000.xlsx")
        assert replaced.read_bytes() == b"damaged"
        assert replaced.name == "Job Search - jane-before-restore-20261003-080000.xlsx"
        assert [r["Job ID"] for r in created.get_records("Jobs")] == ["1", "2"]
        assert sorted(p.name for p in created.path.parent.iterdir() if p.is_file()) == [
            created.path.name
        ]

    def test_restore_without_backups_is_explained(self, created):
        with pytest.raises(WorkbookError, match="no backups"):
            created.restore_latest_backup()

    def test_restore_refuses_while_the_file_is_open(self, created):
        created.backup()
        created.path.with_name(f"~${created.path.name}").write_bytes(b"")
        before = snapshot(created.path)
        with pytest.raises(WorkbookLocked):
            created.restore_latest_backup()
        assert snapshot(created.path) == before

    def test_backing_up_a_missing_file_is_explained(self, wb):
        with pytest.raises(WorkbookMissing):
            wb.backup()


class TestOpenInApp:
    def test_windows_uses_startfile(self, created, monkeypatch):
        opened = []
        monkeypatch.setattr(sys, "platform", "win32")
        monkeypatch.setattr(os, "startfile", opened.append, raising=False)
        created.open_in_app()
        assert opened == [created.path]

    @pytest.mark.parametrize(("platform", "opener"), [("darwin", "open"), ("linux", "xdg-open")])
    def test_other_systems_use_their_opener(self, created, monkeypatch, platform, opener):
        calls = []
        monkeypatch.setattr(sys, "platform", platform)
        monkeypatch.setattr(subprocess, "run", lambda cmd, check: calls.append(cmd))
        created.open_in_app()
        assert calls == [[opener, str(created.path)]]

    def test_a_failure_says_to_install_a_spreadsheet_program(self, created, monkeypatch):
        def fail(cmd, check):
            raise FileNotFoundError(cmd[0])

        monkeypatch.setattr(sys, "platform", "linux")
        monkeypatch.setattr(subprocess, "run", fail)
        with pytest.raises(WorkbookError, match="LibreOffice"):
            created.open_in_app()

    def test_opening_a_missing_file_is_explained(self, wb):
        with pytest.raises(WorkbookMissing):
            wb.open_in_app()


class TestRowsTypedByHand:
    """A row the person typed has no Job ID, but it is theirs and must never be overwritten."""

    def typed_row(self, created):
        book = raw(created)
        ws = book["Jobs"]
        ws["D3"], ws["E3"] = "Applied", "typed by me"
        book.save(created.path)

    def test_an_append_goes_below_it(self, created):
        created.append_rows("Jobs", [job("1")], "Job ID")
        self.typed_row(created)

        created.append_rows("Jobs", [job("2", Notes="agent")], "Job ID")

        ws = raw(created)["Jobs"]
        assert (ws["A3"].value, ws["D3"].value, ws["E3"].value) == (None, "Applied", "typed by me")
        assert (ws["A4"].value, ws["E4"].value) == ("2", "agent")

    def test_an_upsert_that_appends_goes_below_it(self, created):
        created.append_rows("Jobs", [job("1")], "Job ID")
        self.typed_row(created)

        created.upsert_by_key("Jobs", "Job ID", [{"Job ID": "9", "Notes": "agent"}], True)

        ws = raw(created)["Jobs"]
        assert (ws["D3"].value, ws["E3"].value) == ("Applied", "typed by me")
        assert (ws["A4"].value, ws["E4"].value) == ("9", "agent")

    def test_the_autofilter_covers_the_typed_row(self, created):
        created.append_rows("Jobs", [job("1")], "Job ID")
        self.typed_row(created)
        created.append_rows("Jobs", [job("2")], "Job ID")
        assert raw(created)["Jobs"].auto_filter.ref == "A1:E4"

    def test_a_keyless_row_is_not_matched_or_changed_by_an_update(self, created):
        self.typed_row(created)
        assert created.upsert_by_key("Jobs", "Job ID", [{"Job ID": "x", "Notes": "n"}], False) == 0
        assert raw(created)["Jobs"]["E3"].value == "typed by me"

    def test_blank_values_of_an_appended_row_create_no_cells(self, created):
        created.append_rows("Jobs", [job("1")], "Job ID")
        ws = raw(created)["Jobs"]
        assert [c.value for c in ws[2]] == ["1", None, None, "New", None]


class TestTakingTurns:
    """Several agent commands may start together; only one works on the file at a time."""

    @pytest.fixture
    def quick(self, monkeypatch):
        monkeypatch.setattr(workbook_module, "LOCK_WAIT_SECONDS", 0.3)
        monkeypatch.setattr(workbook_module, "LOCK_POLL_SECONDS", 0.02)

    def lock_path(self, wb):
        return wb.path.with_name(f".{wb.path.stem}.agent-lock")

    def test_the_lock_is_released_after_a_session_and_after_a_failed_one(self, created):
        created.get_values("Jobs")
        assert not self.lock_path(created).exists()
        with pytest.raises(RuntimeError), created.session():
            assert self.lock_path(created).exists()
            raise RuntimeError("boom")
        assert not self.lock_path(created).exists()

    def test_a_command_that_finds_it_taken_gives_up_with_a_plain_message(self, created, quick):
        self.lock_path(created).write_text(f"{os.getpid()} 1", encoding="utf-8")
        before = snapshot(created.path)

        with pytest.raises(WorkbookBusy) as error:
            created.append_rows("Jobs", [job("1")], "Job ID")

        message = str(error.value)
        assert "Another update" in message
        assert "run the same command again" in message
        assert not message.startswith(LOCKED_PREFIX)
        assert snapshot(created.path) == before
        assert self.lock_path(created).exists()  # not ours, so not removed

    def test_a_command_waits_for_the_other_to_finish(self, created, monkeypatch):
        monkeypatch.setattr(workbook_module, "LOCK_WAIT_SECONDS", 5)
        monkeypatch.setattr(workbook_module, "LOCK_POLL_SECONDS", 0.02)
        lock = self.lock_path(created)
        lock.write_text(f"{os.getpid()} 1", encoding="utf-8")
        threading.Timer(0.3, lock.unlink).start()

        assert created.append_rows("Jobs", [job("1")], "Job ID") == 1

    def test_a_lock_left_by_a_process_that_died_is_reclaimed(self, created, quick):
        dead = subprocess.Popen([sys.executable, "-c", "pass"])
        dead.wait()
        self.lock_path(created).write_text(f"{dead.pid} 1", encoding="utf-8")

        assert created.append_rows("Jobs", [job("1")], "Job ID") == 1
        assert not self.lock_path(created).exists()

    def test_a_lock_that_is_too_old_is_reclaimed_even_if_its_owner_seems_alive(
        self, created, quick
    ):
        lock = self.lock_path(created)
        lock.write_text(f"{os.getpid()} 1", encoding="utf-8")
        long_ago = time.time() - workbook_module.STALE_LOCK_SECONDS - 60
        os.utime(lock, (long_ago, long_ago))

        assert created.append_rows("Jobs", [job("1")], "Job ID") == 1

    def test_process_liveness(self):
        finished = subprocess.Popen([sys.executable, "-c", "pass"])
        finished.wait()
        assert workbook_module._process_is_alive(os.getpid())
        assert not workbook_module._process_is_alive(finished.pid)

    def test_several_processes_writing_at_once_all_keep_their_rows(self, created):
        scripts = Path(workbook_module.__file__).parent
        rounds = 6
        workers = [
            subprocess.Popen(
                [
                    sys.executable,
                    "-c",
                    CONCURRENT_WORKER,
                    str(scripts),
                    str(created.path),
                    str(created.backups_dir),
                    name,
                    str(rounds),
                ],
                stderr=subprocess.PIPE,
                text=True,
            )
            for name in ("a", "b", "c")
        ]
        errors = [worker.communicate(timeout=120)[1] for worker in workers]

        assert [worker.returncode for worker in workers] == [0, 0, 0], errors
        ids = [r["Job ID"] for r in created.get_records("Jobs")]
        assert sorted(ids) == sorted(f"{n}-{i}" for n in "abc" for i in range(rounds))

    def test_a_second_handle_in_the_same_process_waits_too(self, created, quick):
        other = Workbook(created.path, created.backups_dir)
        with created.session(), pytest.raises(WorkbookBusy):
            other.get_values("Jobs")


CONCURRENT_WORKER = """
import sys
from pathlib import Path
sys.path.insert(0, sys.argv[1])
from workbook import Workbook
book = Workbook(Path(sys.argv[2]), Path(sys.argv[3]))
name = sys.argv[4]
for i in range(int(sys.argv[5])):
    book.upsert_by_key("Jobs", "Job ID", [{"Job ID": f"{name}-{i}", "Notes": name}], True)
"""


class TestSaveByThePersonDuringAnUpdate:
    def test_the_update_is_abandoned_and_their_save_is_untouched(self, created):
        created.append_rows("Jobs", [job("1")], "Job ID")

        with pytest.raises(WorkbookLocked) as error, created.session():
            created.upsert_by_key("Jobs", "Job ID", [{"Job ID": "1", "Score": 5}], False)
            book = raw(created)  # the person saves in Excel meanwhile
            book["Jobs"]["E2"] = "typed in Excel"
            book.save(created.path)

        message = str(error.value)
        assert message.startswith(LOCKED_PREFIX)
        assert "changed while I was updating it" in message
        ws = raw(created)["Jobs"]
        assert ws["E2"].value == "typed in Excel"
        assert ws["B2"].value is None  # the agent's score was not written
        assert [p.name for p in created.path.parent.iterdir() if p.name != "backups"] == [
            created.path.name
        ]

    def test_running_the_same_update_again_works(self, created):
        created.append_rows("Jobs", [job("1")], "Job ID")
        with pytest.raises(WorkbookLocked), created.session():
            created.upsert_by_key("Jobs", "Job ID", [{"Job ID": "1", "Score": 5}], False)
            created.path.write_bytes(created.path.read_bytes() + b"x")  # any change at all

        created.path.write_bytes(created.path.read_bytes()[:-1])
        created.upsert_by_key("Jobs", "Job ID", [{"Job ID": "1", "Score": 5}], False)
        assert created.get_records("Jobs")[0]["Score"] == "5"

    def test_a_file_that_appeared_meanwhile_is_not_overwritten(self, wb):
        with pytest.raises(WorkbookLocked), wb.session(create=True):
            wb.ensure_tabs(SPECS)
            wb.path.write_bytes(b"someone else made this")
        assert wb.path.read_bytes() == b"someone else made this"


class TestReadOnlyFile:
    @pytest.fixture
    def read_only(self, created):
        created.path.chmod(stat.S_IREAD)
        yield created
        created.path.chmod(stat.S_IWRITE | stat.S_IREAD)

    def test_writing_says_how_to_clear_the_read_only_flag(self, read_only):
        before = snapshot(read_only.path)
        with pytest.raises(WorkbookReadOnly) as error:
            read_only.append_rows("Jobs", [job("1")], "Job ID")

        message = str(error.value)
        assert "read-only" in message
        assert "Properties" in message
        assert not message.startswith(LOCKED_PREFIX)
        assert snapshot(read_only.path) == before

    def test_it_is_not_mistaken_for_an_open_file_and_can_still_be_read(self, read_only):
        assert read_only.is_read_only()
        assert not read_only.is_open()
        assert read_only.get_headers("Jobs") == JOB_HEADERS
        assert read_only.clear_lock_files() == []

    def test_require_closed_reports_it_first(self, read_only):
        read_only.path.with_name(f"~${read_only.path.name}").write_bytes(b"")
        with pytest.raises(WorkbookReadOnly):
            read_only.require_closed()

    def test_a_normal_file_is_not_read_only(self, created):
        assert not created.is_read_only()


class TestExcelTables:
    def add_table(self, wb, ref="A1:E3", sheet_filter=True):
        book = raw(wb)
        ws = book["Jobs"]
        table = Table(displayName="JobsTable", ref=ref)
        table.autoFilter = AutoFilter(ref=ref)
        ws.add_table(table)
        ws.auto_filter.ref = ref if sheet_filter else None
        book.save(wb.path)

    def test_new_rows_join_the_table_and_there_is_no_sheet_filter_over_it(self, created):
        created.append_rows("Jobs", [job("1"), job("2")], "Job ID")
        self.add_table(created)

        created.append_rows("Jobs", [job("3"), job("4")], "Job ID")

        ws = raw(created)["Jobs"]
        assert ws.auto_filter.ref is None
        table = ws.tables["JobsTable"]
        assert table.ref == "A1:E5"
        assert table.autoFilter.ref == "A1:E5"
        assert ws["A5"].value == "4"

    def test_a_table_is_never_shrunk_below_the_data_it_has(self, created):
        created.append_rows("Jobs", [job("1")], "Job ID")
        self.add_table(created, ref="A1:E10")
        created.append_rows("Jobs", [job("2")], "Job ID")
        assert raw(created)["Jobs"].tables["JobsTable"].ref == "A1:E10"

    def test_a_table_with_a_totals_row_is_left_alone(self, created):
        created.append_rows("Jobs", [job("1")], "Job ID")
        book = raw(created)
        table = Table(displayName="JobsTable", ref="A1:E3", totalsRowCount=1)
        book["Jobs"].add_table(table)
        book.save(created.path)

        created.upsert_by_key("Jobs", "Job ID", [{"Job ID": "1", "Notes": "n"}], False)

        assert raw(created)["Jobs"].tables["JobsTable"].ref == "A1:E3"

    def test_a_sheet_without_a_table_still_gets_its_filter(self, created):
        created.append_rows("Jobs", [job("1")], "Job ID")
        assert raw(created)["Jobs"].auto_filter.ref == "A1:E2"


class TestListsExportedFromGoogle:
    def exported(self, wb, hidden_by_person=()):
        """What File > Download > Excel gives: every tab visible, in whatever shape."""
        book = openpyxl.Workbook()
        book.active.title = "Jobs"
        book["Jobs"].append(JOB_HEADERS)
        for name in ("Rules", "Hidden"):
            book.create_sheet(name).append(
                ["Rule", "Type"] if name == "Rules" else ["Key", "Value"]
            )
        for name in hidden_by_person:
            book[name].sheet_state = "hidden"
        book.save(wb.path)

    def test_bookkeeping_tabs_become_hidden_and_others_stay_visible(self, wb):
        self.exported(wb)
        wb.ensure_tabs(SPECS)
        states = {ws.title: ws.sheet_state for ws in raw(wb)}
        assert states == {
            "Jobs": "visible",
            "Rules": "visible",
            "Hidden": "hidden",
            "Prefs": "visible",
        }

    def test_a_tab_the_person_hid_is_never_shown_again(self, wb):
        self.exported(wb, hidden_by_person=("Rules",))
        wb.ensure_tabs(SPECS)
        assert raw(wb)["Rules"].sheet_state == "hidden"

    def test_nothing_is_saved_when_everything_is_already_right(self, created):
        before = snapshot(created.path)
        created.ensure_tabs(SPECS)
        assert snapshot(created.path) == before


class TestHyperlinkStyle:
    def test_a_cleared_link_loses_its_blue_underlined_style(self, created):
        created.append_rows("Jobs", [job("1", Link="https://a.b/c")], "Job ID")
        assert raw(created)["Jobs"]["C2"].style == "Hyperlink"

        created.upsert_by_key("Jobs", "Job ID", [{"Job ID": "1", "Link": ""}], False)

        cell = raw(created)["Jobs"]["C2"]
        assert cell.hyperlink is None
        assert cell.style != "Hyperlink"

    def test_text_that_is_not_a_link_does_not_touch_other_styles(self, created):
        created.append_rows("Jobs", [job("1", Link="not a link")], "Job ID")
        assert raw(created)["Jobs"]["C2"].style == "Normal"


class TestLeftovers:
    def test_leftover_temp_files_of_both_naming_schemes_are_found(self, created):
        stem = created.path.stem
        mine = [
            created.path.with_name(f".{stem}.abc123.saving.xlsx"),
            created.path.with_name(f".{stem}.saving.xlsx"),
        ]
        for leftover in (*mine, created.path.with_name(".other.abc.saving.xlsx")):
            leftover.write_bytes(b"half")
        assert created.leftover_temp_files() == sorted(mine)

    def test_clear_lock_files_also_removes_leftovers_and_a_dead_updates_lock(self, created):
        dead = subprocess.Popen([sys.executable, "-c", "pass"])
        dead.wait()
        lock = created.path.with_name(f".{created.path.stem}.agent-lock")
        lock.write_text(f"{dead.pid} 1", encoding="utf-8")
        leftover = created.path.with_name(f".{created.path.stem}.zzz.saving.xlsx")
        leftover.write_bytes(b"half")

        removed = created.clear_lock_files()

        assert sorted(removed) == sorted([lock, leftover])
        assert sorted(p.name for p in created.path.parent.iterdir()) == [created.path.name]

    def test_an_update_running_right_now_is_left_alone(self, created):
        lock = created.path.with_name(f".{created.path.stem}.agent-lock")
        lock.write_text(f"{os.getpid()} 1", encoding="utf-8")
        leftover = created.path.with_name(f".{created.path.stem}.zzz.saving.xlsx")
        leftover.write_bytes(b"in progress")

        assert created.clear_lock_files() == []
        assert lock.exists()
        assert leftover.exists()


class TestPlainErrors:
    def test_a_backup_that_cannot_be_written_is_explained(self, created):
        created.backups_dir.write_text("a file where the folder should be", encoding="utf-8")
        with pytest.raises(WorkbookError, match="Couldn't make a backup"):
            created.backup()

    def test_a_restore_that_cannot_copy_leaves_the_file_alone(self, created, monkeypatch):
        created.backup()
        before = snapshot(created.path)

        def fail(*args):
            raise OSError("disk full")

        monkeypatch.setattr(workbook_module.shutil, "copyfile", fail)
        with pytest.raises(WorkbookError, match="disk full"):
            created.restore_latest_backup()
        assert snapshot(created.path) == before

    def test_a_save_that_fails_inside_openpyxl_is_explained_and_cleaned_up(
        self, created, monkeypatch
    ):
        before = snapshot(created.path)

        def explode(self, filename):
            raise RuntimeError("zip trouble")

        monkeypatch.setattr(openpyxl.Workbook, "save", explode)
        with pytest.raises(WorkbookError, match="Couldn't save.*zip trouble"):
            created.append_rows("Jobs", [job("1")], "Job ID")

        assert snapshot(created.path) == before
        assert [p.name for p in created.path.parent.iterdir() if p.name != "backups"] == [
            created.path.name
        ]

    def test_a_temp_file_that_vanishes_before_the_replace_is_explained(self, created, monkeypatch):
        def vanish(source, target):
            raise FileNotFoundError(2, "No such file", str(source))

        monkeypatch.setattr(workbook_module.os, "replace", vanish)
        with pytest.raises(WorkbookError, match="Couldn't save"):
            created.append_rows("Jobs", [job("1")], "Job ID")

    def test_messages_name_the_real_command_when_the_owner_gives_it(self, tmp_path):
        command = "uv run python scripts/sheet.py --profile jane"
        owned = Workbook(tmp_path / "a.xlsx", tmp_path / "backups", command=command)
        owned.path.write_bytes(b"junk")
        with pytest.raises(WorkbookCorrupt) as error:
            owned.verify()
        assert f"`{command} restore-backup`" in str(error.value)
        assert "<name>" not in str(error.value)

    def test_the_default_command_is_a_template(self, tmp_path):
        plain = Workbook(tmp_path / "a.xlsx", tmp_path / "backups")
        plain.path.write_bytes(b"junk")
        with pytest.raises(WorkbookCorrupt, match="--profile <name> restore-backup"):
            plain.verify()

    def test_the_damaged_message_covers_password_protection(self, wb):
        wb.path.write_bytes(b"junk")
        with pytest.raises(WorkbookCorrupt, match="password-protected"):
            wb.verify()
