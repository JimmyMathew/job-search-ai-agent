"""Storage layer: one Excel workbook (.xlsx) on the person's own computer.

This module knows files, tabs, headers and cells. It knows nothing about jobs; the job-list
layout and rules live in sheet_schema.py and sheet.py.

The person opens this file in Excel while the agent also reads and writes it, so the hard
parts here are about not losing their data:

* A write never happens while the file is open elsewhere (Excel locks it), and a save goes
  through a temporary file that replaces the original in one step, so the original survives
  any failure. If the person saved the file while an update was running, the update is
  abandoned rather than overwriting their save. Reading is always allowed.
* Only one agent process works on the file at a time (a lock file held from load to save), so
  two commands run in parallel can't lose each other's changes.
* A file that can't be loaded is reported, never replaced by a fresh one.
* Cells are written safely: text starting with "=" stays text, illegal characters are
  stripped, very long text is cut to what Excel can hold.
* A copy is kept before each run (`backup`) so a damaged file can be put back.

Several operations in one command share a single load and a single save through `session()`.
"""

from __future__ import annotations

import contextlib
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from collections.abc import Iterator, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
from config import UserFacingError
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet

# Scripts and /job-run look for this prefix to know the file is open and should be retried.
LOCKED_PREFIX = "WORKBOOK_OPEN:"

LINK_HEADER = "Link"
HYPERLINK_STYLE = "Hyperlink"
MAX_CELL_CHARS = 32_000  # Excel's hard limit is 32,767 characters per cell
MAX_LINK_CHARS = 2_000  # Excel ignores hyperlinks of about 2,080 characters or more
MAX_LIST_VALIDATION_CHARS = 255  # Excel's limit for a typed-in dropdown list
LAST_SHEET_ROW = 1_048_576
FREEZE_BELOW_HEADER = "A2"

MAX_BACKUPS = 10
BACKUP_STAMP_FORMAT = "%Y%m%d-%H%M%S"
REPLACE_RETRIES = 3  # antivirus and sync tools briefly hold a freshly written file
REPLACE_RETRY_SECONDS = 0.5

LOCK_WAIT_SECONDS = 30  # how long a command waits for another command to finish with the file
LOCK_POLL_SECONDS = 0.1
LOCK_RELEASE_TRIES = 50  # a waiter peeking at the lock file briefly blocks deleting it on Windows
STALE_LOCK_SECONDS = 600  # a lock this old is left over from a crash, whatever its owner says
WINDOWS_STILL_ACTIVE = 259
WINDOWS_ACCESS_DENIED = 5
WINDOWS_QUERY_LIMITED_INFORMATION = 0x1000

HEADER_FILL_COLOR = "D9E1F2"
MIN_COLUMN_WIDTH = 10
COLUMN_WIDTH_PADDING = 2

# What messages tell the person to run, until the owner of the Workbook says which profile it is.
DEFAULT_COMMAND = "uv run python scripts/sheet.py --profile <name>"

_LINK_RE = re.compile(r"https?://\S+")


class WorkbookError(UserFacingError):
    """A problem with the job-list file that the person can fix."""


class WorkbookLocked(WorkbookError):
    """The file is open in another program, so it can't be changed right now."""


class WorkbookBusy(WorkbookError):
    """Another agent command is still working on the file; try again in a moment."""


class WorkbookReadOnly(WorkbookError):
    """The file is marked read-only, so it can't be changed."""


class WorkbookMissing(WorkbookError):
    """There is no job-list file where one is expected."""


class WorkbookCorrupt(WorkbookError):
    """The file exists but can't be read as a workbook."""


@dataclass(frozen=True)
class TabSpec:
    """One tab a workbook should have: its headers, example rows for a brand-new tab, whether
    it is hidden, and column widths (in characters) by header."""

    name: str
    headers: Sequence[str]
    seed_rows: Sequence[Sequence[str]] = ()
    hidden: bool = False
    widths: Mapping[str, int] = field(default_factory=dict)


# ---- Pure helpers ---------------------------------------------------------------------


def normalise(value: object) -> str:
    """A cell's value as plain text: blank for empty, dates as YYYY-MM-DD, whole numbers
    without the ".0" Excel's number storage adds."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def storable(value: object) -> int | float | str | None:
    """The value to put in a cell. Numbers stay numbers (so Score sorts properly); everything
    else becomes text with illegal characters removed and a length Excel can hold. Blank
    becomes an empty cell."""
    if isinstance(value, int | float) and not isinstance(value, bool):
        return value
    text = ILLEGAL_CHARACTERS_RE.sub("", normalise(value))[:MAX_CELL_CHARS]
    return text or None


def is_linkable(url: str) -> bool:
    """A web address safe to turn into a clickable link (Excel drops very long ones)."""
    return len(url) < MAX_LINK_CHARS and _LINK_RE.fullmatch(url) is not None


def _required_key(row: Mapping[str, Any], header: str) -> str:
    key = normalise(row.get(header)).strip()
    if not key:
        raise WorkbookError(
            f"A row to write has no {header!r}, so it can't be matched: {dict(row)}"
        )
    return key


def _process_is_alive(pid: int) -> bool:
    """Whether a process with this ID is running. (os.kill can't be used on Windows: any signal
    but Ctrl-C there terminates the process.)"""
    if os.name == "nt":
        import ctypes

        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        handle = kernel32.OpenProcess(WINDOWS_QUERY_LIMITED_INFORMATION, False, pid)
        if not handle:
            return ctypes.get_last_error() == WINDOWS_ACCESS_DENIED  # exists, but not ours
        try:
            exit_code = ctypes.c_ulong()
            kernel32.GetExitCodeProcess(handle, ctypes.byref(exit_code))
            return exit_code.value == WINDOWS_STILL_ACTIVE
        finally:
            kernel32.CloseHandle(handle)
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    except PermissionError:
        return True
    return True


def _lock_owner(text: str) -> int | None:
    """The process ID a lock file says it belongs to."""
    first = text.split(maxsplit=1)[:1]
    return int(first[0]) if first and first[0].isdigit() else None


def _lock_is_stale(lock: Path) -> bool:
    """A lock nobody is working under any more: too old, or its owner no longer runs."""
    try:
        age = time.time() - lock.stat().st_mtime
        owner = _lock_owner(lock.read_text(encoding="utf-8"))
    except OSError:
        return False  # it just went away, which is not a reason to wait or to delete anything
    return age > STALE_LOCK_SECONDS or (owner is not None and not _process_is_alive(owner))


def _release_lock(lock: Path, token: str) -> None:
    """Delete our lock file, unless a stale-lock reclaimer already replaced it with theirs. On
    Windows deleting fails for a moment while a waiting command is reading the file, and a lock
    left behind would block this very process's next update, so try again."""
    for _ in range(LOCK_RELEASE_TRIES):
        try:
            if lock.read_text(encoding="utf-8") == token:
                lock.unlink()
            return
        except FileNotFoundError:
            return
        except OSError:
            time.sleep(LOCK_POLL_SECONDS / 5)


# ---- Worksheet helpers (no file access) -----------------------------------------------------


def _header_cells(ws: Worksheet) -> list[str]:
    """The header row as text, without trailing blank cells."""
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = [normalise(value).strip() for value in row]
    while headers and not headers[-1]:
        headers.pop()
    return headers


def _columns(ws: Worksheet) -> dict[str, int]:
    """Header name -> 1-based column. Rows are always placed by these names, never by
    position, so a column the person moved, added or deleted can't misalign anything."""
    columns: dict[str, int] = {}
    for index, header in enumerate(_header_cells(ws), start=1):
        if header:
            columns.setdefault(header, index)
    return columns


def _key_column(columns: Mapping[str, int], tab: str, header: str, command: str) -> int:
    if header not in columns:
        raise WorkbookError(
            f"The '{tab}' tab has no '{header}' column (it may have been deleted or renamed, "
            f"and the agent needs it). Running `{command} init` adds it back."
        )
    return columns[header]


def _last_data_row(ws: Worksheet) -> int:
    """The last row with anything in any cell, so new rows never land on a row the person typed
    by hand (which has no Job ID). Not ws.max_row, which also counts rows that are merely
    formatted."""
    for row in range(ws.max_row, 1, -1):
        if any(normalise(cell.value).strip() for cell in ws[row]):
            return row
    return 1


def _rows_by_key(ws: Worksheet, key_col: int) -> dict[str, int]:
    """Key -> row number, first occurrence winning; read now, so a re-sorted sheet is safe."""
    rows: dict[str, int] = {}
    cells = ws.iter_rows(min_row=2, min_col=key_col, max_col=key_col, values_only=True)
    for row_number, (value,) in enumerate(cells, start=2):
        key = normalise(value).strip()
        if key:
            rows.setdefault(key, row_number)
    return rows


def _write_cell(cell: Cell, header: str, value: object) -> None:
    stored = storable(value)
    cell.value = stored
    if isinstance(stored, str):
        cell.data_type = "s"  # assigning "=..." makes a formula; force it back to plain text
    if header == LINK_HEADER:
        _set_link(cell, stored)


def _set_link(cell: Cell, stored: object) -> None:
    if isinstance(stored, str) and is_linkable(stored):
        cell.hyperlink = stored
        cell.style = HYPERLINK_STYLE
    else:
        cell.hyperlink = None
        if cell.style == HYPERLINK_STYLE:  # no longer a link, so no longer blue and underlined
            cell.style = "Normal"


def _write_row(
    ws: Worksheet,
    columns: Mapping[str, int],
    row_number: int,
    values: Mapping[str, Any],
    skip_blank: bool = False,
) -> None:
    """Write the values whose header the sheet has. A blank clears the cell it names, unless
    `skip_blank` (a brand-new row has nothing to clear)."""
    for header, value in values.items():
        if header in columns and not (skip_blank and storable(value) is None):
            _write_cell(ws.cell(row_number, columns[header]), header, value)


def _column_width(spec: TabSpec, header: str) -> int:
    """The width asked for in the spec, else one that fits the header's title."""
    return spec.widths.get(header) or max(MIN_COLUMN_WIDTH, len(header) + COLUMN_WIDTH_PADDING)


def _write_header(ws: Worksheet, column: int, spec: TabSpec, header: str) -> None:
    cell = ws.cell(1, column, header)
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor=HEADER_FILL_COLOR)
    ws.column_dimensions[get_column_letter(column)].width = _column_width(spec, header)


def _extend_table(table: Table, last_row: int) -> None:
    """Grow an Excel table (Format as Table) to the last data row, so new rows join it."""
    if table.totalsRowCount:
        return  # a totals row sits below the data; growing the table would swallow it
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    last = max(max_row, last_row, min_row + 1)  # never shrink; a table needs a data row
    table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{last}"
    if table.autoFilter is not None:
        table.autoFilter.ref = table.ref


def _refresh_ranges(ws: Worksheet) -> None:
    """Make the filter arrows (or the person's table) cover exactly the data, however many rows
    were added. A sheet-level filter over a table makes Excel call the file damaged, so a
    sheet that has a table gets no sheet-level filter."""
    headers = _header_cells(ws)
    if not headers:
        return
    last_row = _last_data_row(ws)
    if ws.tables:
        ws.auto_filter.ref = None
        for table in ws.tables.values():
            _extend_table(table, last_row)
    else:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"


def _covers_column(validation: DataValidation, column: int) -> bool:
    return any(r.min_col <= column <= r.max_col for r in validation.sqref.ranges)


def _is_list_validation(validation: DataValidation, formula: str, cells: str) -> bool:
    return (
        validation.type == "list"
        and validation.formula1 == formula
        and str(validation.sqref) == cells
        and bool(validation.allow_blank)
        and not validation.showDropDown
    )


# ---- The workbook -----------------------------------------------------------------------------


class Workbook:
    """One .xlsx file plus the folder its automatic backups go to. `command` is what messages
    tell the person to run to repair things (see config.sheet_command)."""

    def __init__(self, path: Path, backups_dir: Path, command: str = DEFAULT_COMMAND) -> None:
        self.path = path
        self.backups_dir = backups_dir
        self.command = command
        self._book: openpyxl.Workbook | None = None
        self._changed = False
        self._loaded_stat: tuple[int, int] | None = None

    # ---- State of the file ---------------------------------------------------------

    def exists(self) -> bool:
        return self.path.is_file()

    def is_read_only(self) -> bool:
        return self.exists() and not os.access(self.path, os.W_OK)

    def lock_files(self) -> list[Path]:
        """Lock files that Excel (and LibreOffice) leave beside a file while it is open.
        Excel shortens the name of a long file in its lock file, so both spellings count."""
        name = self.path.name
        candidates = {f"~${name}", f"~${name[2:]}", f".~lock.{name}#"}
        return [
            lock for lock in (self.path.with_name(c) for c in sorted(candidates)) if lock.exists()
        ]

    def leftover_temp_files(self) -> list[Path]:
        """Half-written copies a save that was killed leaves beside the workbook."""
        pattern = re.compile(re.escape(f".{self.path.stem}.") + r"(?:.+\.)?saving\.xlsx")
        if not self.path.parent.is_dir():
            return []
        return sorted(p for p in self.path.parent.iterdir() if pattern.fullmatch(p.name))

    def _held_by_another_program(self) -> bool:
        """Windows refuses write access to a file another program has open; a lock file can be
        left behind by a crash, but this check can't be fooled that way. Elsewhere there is
        no such check, so only lock files count. A read-only file refuses write access too,
        which says nothing about another program."""
        if os.name != "nt" or not self.exists() or self.is_read_only():
            return False
        try:
            with self.path.open("r+b"):
                return False
        except PermissionError:
            return True

    def is_open(self) -> bool:
        """Whether the file is open in a spreadsheet program (and so must not be written)."""
        return bool(self.lock_files()) or self._held_by_another_program()

    def require_closed(self) -> None:
        """Raise unless the file can be written now: not read-only, not open elsewhere."""
        if self.is_read_only():
            raise WorkbookReadOnly(
                f'Your job list "{self.path.name}" is marked read-only, so it can\'t be updated. '
                "Right-click the file, choose Properties, untick Read-only and click OK "
                "(on a Mac: Get Info, untick Locked), then run the command again."
            )
        if self.is_open():
            raise WorkbookLocked(self._locked_message())

    def clear_lock_files(self) -> list[Path]:
        """Delete what a crash leaves behind (Excel's lock files, half-written copies, a dead
        update's lock) and return it. Refuses while Windows says the file really is open, and
        leaves the files of an update that is running right now alone."""
        if self._held_by_another_program():
            raise WorkbookLocked(self._locked_message())
        removed = self.lock_files()
        agent_lock = self._agent_lock_path()
        if not agent_lock.exists() or _lock_is_stale(agent_lock):
            removed += [agent_lock] if agent_lock.exists() else []
            removed += self.leftover_temp_files()
        for leftover in removed:
            leftover.unlink(missing_ok=True)
        return removed

    def _locked_message(self) -> str:
        return (
            f'{LOCKED_PREFIX} Your job list "{self.path.name}" is open in another program '
            "(usually Excel), so it can't be updated right now. Please save it, close it, "
            "and tell me when it is closed."
        )

    # ---- Taking turns with other agent commands ------------------------------------

    def _agent_lock_path(self) -> Path:
        return self.path.with_name(f".{self.path.stem}.agent-lock")

    @contextlib.contextmanager
    def _exclusive_access(self) -> Iterator[None]:
        """Hold the agent lock: only one command at a time may load, change and save the file.
        Two commands started together (Claude often runs several at once) would otherwise each
        save their own copy and one would lose its changes."""
        lock = self._agent_lock_path()
        token = self._acquire(lock)
        try:
            yield
        finally:
            _release_lock(lock, token)

    def _acquire(self, lock: Path) -> str:
        """Create the lock file (atomically: it either didn't exist or we fail), waiting for
        another command to finish and reclaiming a lock a crashed command left behind."""
        token = f"{os.getpid()} {time.time_ns()}"
        deadline = time.monotonic() + LOCK_WAIT_SECONDS
        while True:
            try:
                fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            except (FileExistsError, PermissionError) as exc:
                if isinstance(exc, PermissionError) and not os.access(lock.parent, os.W_OK):
                    raise WorkbookError(
                        f"The folder {lock.parent} is read-only, so your job list can't be updated."
                    ) from exc
                self._wait_for_lock(lock, deadline)
            except OSError as exc:
                raise WorkbookError(f"Couldn't start updating your job list: {exc}") from exc
            else:
                with os.fdopen(fd, "w", encoding="utf-8") as handle:
                    handle.write(token)
                return token

    @staticmethod
    def _wait_for_lock(lock: Path, deadline: float) -> None:
        """Wait a moment for the lock to clear, remove it if it's left over, or give up."""
        try:
            seen = lock.read_text(encoding="utf-8")
        except OSError:
            return  # released while we looked: try to take it straight away
        if _lock_is_stale(lock):
            with contextlib.suppress(OSError):
                if lock.read_text(encoding="utf-8") == seen:  # still the one we judged stale
                    lock.unlink()
            return
        if time.monotonic() >= deadline:
            raise WorkbookBusy(
                "Another update of your job list is still running. Wait a few seconds and "
                "run the same command again."
            )
        time.sleep(LOCK_POLL_SECONDS)

    # ---- Sessions: one load, one save ----------------------------------------------

    @contextlib.contextmanager
    def session(self, *, create: bool = False) -> Iterator[None]:
        """Load the file once, let several operations share it, and save once at the end if
        anything changed. Nested sessions join the outer one. Nothing is saved if an error
        interrupts the session. The whole session holds the agent lock. `create` lets the
        session start a new file when none exists."""
        if self._book is not None:
            yield
            return
        if create:
            self._make_folder()
        if not self.path.parent.is_dir():
            raise WorkbookMissing(f"There is no job list at {self.path}.")
        with self._exclusive_access():
            self._changed = False
            self._loaded_stat = self._stat()
            self._book = self._load() if self.exists() else self._new_book(create)
            try:
                yield
                if self._changed:
                    self._save(self._book)
            finally:
                self._book = None
                self._changed = False

    def _make_folder(self) -> None:
        try:
            self.path.parent.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            raise WorkbookError(f"Couldn't create the folder {self.path.parent}: {exc}") from exc

    def _stat(self) -> tuple[int, int] | None:
        """What the file looks like now, to notice if someone else saved it meanwhile."""
        try:
            info = self.path.stat()
        except OSError:
            return None
        return (info.st_mtime_ns, info.st_size)

    def _require_book(self) -> openpyxl.Workbook:
        if self._book is None:
            raise RuntimeError("Workbook used outside a session")
        return self._book

    def _begin_change(self) -> None:
        """Called before an operation first changes anything: refuses while the file is open,
        so a locked file is reported before any work is done."""
        if not self._changed:
            self.require_closed()
            self._changed = True

    def _load(self) -> openpyxl.Workbook:
        try:
            return openpyxl.load_workbook(self.path)
        except OSError as exc:
            raise WorkbookError(f"Couldn't read your job list {self.path}: {exc}") from exc
        except Exception as exc:  # openpyxl raises many unrelated types for a damaged file
            raise WorkbookCorrupt(self._corrupt_message()) from exc

    def _new_book(self, create: bool) -> openpyxl.Workbook:
        if not create:
            raise WorkbookMissing(f"There is no job list at {self.path}.")
        self._begin_change()
        book = openpyxl.Workbook()
        book.remove(book.active)  # the empty "Sheet" openpyxl starts with
        return book

    def _corrupt_message(self) -> str:
        return (
            f"Your job list {self.path} can't be opened: it is damaged or password-protected. "
            f"It has not been changed. Automatic copies from before each run are in "
            f"{self.backups_dir}: run `{self.command} restore-backup` to put the newest one "
            "back (the file it replaces is kept there too)."
        )

    def _new_temp_file(self) -> Path:
        """An empty file with a name no other save uses, beside the workbook so the final replace
        is one step on the same disk. Never named "~$...", which Excel would mistake for one of
        its own lock files."""
        try:
            fd, name = tempfile.mkstemp(
                dir=self.path.parent, prefix=f".{self.path.stem}.", suffix=".saving.xlsx"
            )
        except OSError as exc:
            raise WorkbookError(f"Couldn't save your job list {self.path}: {exc}") from exc
        os.close(fd)
        return Path(name)

    def _save(self, book: openpyxl.Workbook) -> None:
        self.require_closed()
        for ws in book.worksheets:
            _refresh_ranges(ws)
        temp = self._new_temp_file()
        try:
            book.save(temp)
            self._refuse_if_changed()
        except WorkbookError:
            temp.unlink(missing_ok=True)
            raise
        except Exception as exc:
            temp.unlink(missing_ok=True)
            raise WorkbookError(f"Couldn't save your job list {self.path}: {exc}") from exc
        self._install(temp)

    def _refuse_if_changed(self) -> None:
        """Replacing the file would erase a save the person made while this update was running
        (typically in Excel, which doesn't always leave a lock file we can see)."""
        if self._stat() != self._loaded_stat:
            raise WorkbookLocked(
                f'{LOCKED_PREFIX} Your job list "{self.path.name}" was changed while I was '
                "updating it (probably saved in Excel). I did not overwrite your changes and "
                "nothing is lost. Please save it, close it, and tell me when it is closed: "
                "I will run the command again."
            )

    def _install(self, temp: Path) -> None:
        """Replace the workbook with a fully written temp file. If that keeps failing, the
        original is left exactly as it was and the temp file is removed."""
        try:
            for attempt in range(REPLACE_RETRIES + 1):
                try:
                    os.replace(temp, self.path)
                    return
                except PermissionError:
                    if attempt < REPLACE_RETRIES:
                        time.sleep(REPLACE_RETRY_SECONDS)
                except OSError as exc:
                    raise WorkbookError(f"Couldn't save your job list {self.path}: {exc}") from exc
            raise WorkbookLocked(self._locked_message())
        finally:
            temp.unlink(missing_ok=True)

    def _sheet(self, tab: str) -> Worksheet:
        book = self._require_book()
        if tab not in book.sheetnames:
            raise WorkbookError(
                f"Your job list has no '{tab}' tab (it may have been deleted or renamed). "
                f"Running `{self.command} init` creates it again."
            )
        return book[tab]

    # ---- Reading -------------------------------------------------------------------

    def verify(self) -> None:
        """Raise WorkbookCorrupt (or WorkbookMissing) unless the file loads as a workbook."""
        with self.session():
            pass

    def tab_names(self) -> list[str]:
        with self.session():
            return list(self._require_book().sheetnames)

    def get_headers(self, tab: str) -> list[str]:
        with self.session():
            return _header_cells(self._sheet(tab))

    def get_values(self, tab: str) -> list[list[str]]:
        """Every non-blank row of a tab as text, header row first."""
        with self.session():
            raw_rows = self._sheet(tab).iter_rows(values_only=True)
            rows = [[normalise(value) for value in row] for row in raw_rows]
            return [row for row in rows if any(row)]

    def get_records(self, tab: str) -> list[dict[str, str]]:
        """Every data row of a tab as {header: text}, like a CSV reader."""
        values = self.get_values(tab)
        if not values:
            return []
        headers, *rows = values
        return [{h: v for h, v in zip(headers, row, strict=False) if h} for row in rows]

    # ---- Writing -------------------------------------------------------------------

    def ensure_tabs(self, specs: Sequence[TabSpec]) -> None:
        """Create the file if needed and add whatever is missing, never overwriting anything.

        Missing tabs are created in order with styled headers, a frozen header row and their
        seed rows. Existing tabs only gain headers they lack, at the end of the header row,
        so data and anything the person typed is untouched; a tab the spec says is hidden is
        hidden (a list exported from the old Google version has them visible), but a hidden
        tab is never shown again. The first tab starts selected in a new file."""
        with self.session(create=True):
            book = self._require_book()
            brand_new = not book.sheetnames
            for spec in specs:
                if spec.name in book.sheetnames:
                    self._add_missing_headers(book[spec.name], spec)
                    self._hide_if_wanted(book[spec.name], spec)
                else:
                    self._create_tab(spec)
            if brand_new:
                self._select_tab(book, specs[0].name)

    def _create_tab(self, spec: TabSpec) -> None:
        self._begin_change()
        ws = self._require_book().create_sheet(spec.name)
        for column, header in enumerate(spec.headers, start=1):
            _write_header(ws, column, spec, header)
        for row_number, seed in enumerate(spec.seed_rows, start=2):
            for column, value in enumerate(seed, start=1):
                _write_cell(ws.cell(row_number, column), spec.headers[column - 1], value)
        ws.freeze_panes = FREEZE_BELOW_HEADER
        if spec.hidden:
            ws.sheet_state = "hidden"

    def _add_missing_headers(self, ws: Worksheet, spec: TabSpec) -> None:
        present = _header_cells(ws)
        missing = [header for header in spec.headers if header not in present]
        if missing:
            self._begin_change()
        for column, header in enumerate(missing, start=len(present) + 1):
            _write_header(ws, column, spec, header)

    def _hide_if_wanted(self, ws: Worksheet, spec: TabSpec) -> None:
        if spec.hidden and ws.sheet_state != "hidden":
            self._begin_change()
            ws.sheet_state = "hidden"

    @staticmethod
    def _select_tab(book: openpyxl.Workbook, name: str) -> None:
        book.active = book.sheetnames.index(name)
        for ws in book.worksheets:  # two selected tabs would make Excel edit both at once
            ws.sheet_view.tabSelected = ws.title == name

    def append_rows(
        self, tab: str, rows: Sequence[Mapping[str, Any]], unique_header: str | None = None
    ) -> int:
        """Add rows after the last row that has anything in it and return how many were added.

        Cells are placed by header name; a key with no matching column is skipped, because
        the person may have deleted that column. With `unique_header`, a row whose value
        there is already in the tab (or earlier in `rows`) is skipped, which makes retrying a
        write harmless."""
        with self.session():
            ws = self._sheet(tab)
            columns = _columns(ws)
            key_col = (
                _key_column(columns, tab, unique_header, self.command) if unique_header else None
            )
            known = set(_rows_by_key(ws, key_col)) if key_col else set()
            next_row = _last_data_row(ws) + 1
            added = 0
            for row in rows:
                if unique_header:
                    key = _required_key(row, unique_header)
                    if key in known:
                        continue
                    known.add(key)
                self._begin_change()
                _write_row(ws, columns, next_row, row, skip_blank=True)
                next_row += 1
                added += 1
            return added

    def upsert_by_key(
        self,
        tab: str,
        key_header: str,
        rows: Sequence[Mapping[str, Any]],
        append_missing: bool,
    ) -> int:
        """Update the row whose `key_header` matches each dict, writing only the headers that
        dict contains (a blank value clears that cell), so every other column, including
        what the person typed, is left alone. Rows with an unknown key are appended if
        `append_missing`, otherwise skipped. Returns how many rows were written.

        Rows are matched by key when this runs, not by a remembered position, so a job list
        the person has sorted or filtered updates the right rows."""
        with self.session():
            ws = self._sheet(tab)
            columns = _columns(ws)
            key_col = _key_column(columns, tab, key_header, self.command)
            row_of_key = _rows_by_key(ws, key_col)
            next_row = _last_data_row(ws) + 1
            written = 0
            for row in rows:
                key = _required_key(row, key_header)
                if key in row_of_key:
                    target = row_of_key[key]
                    values = {h: v for h, v in row.items() if h != key_header}
                elif append_missing:
                    target, values = next_row, row
                    row_of_key[key] = target
                    next_row += 1
                else:
                    continue
                self._begin_change()
                _write_row(ws, columns, target, values)
                written += 1
            return written

    def set_list_validation(self, tab: str, header: str, values: Sequence[str]) -> None:
        """Give one column a dropdown of `values` (below the header). Any dropdown or rule that
        already covers that column is replaced; ones on other columns are kept. Nothing is
        changed (and so nothing saved) when the column already has exactly this dropdown."""
        formula = '"' + ",".join(values) + '"'
        if any("," in v for v in values) or len(formula) > MAX_LIST_VALIDATION_CHARS:
            raise ValueError("Dropdown values must be short and must not contain commas")
        with self.session():
            ws = self._sheet(tab)
            column = _key_column(_columns(ws), tab, header, self.command)
            letter = get_column_letter(column)
            cells = f"{letter}2:{letter}{LAST_SHEET_ROW}"
            existing = ws.data_validations.dataValidation
            covering = [v for v in existing if _covers_column(v, column)]
            if len(covering) == 1 and _is_list_validation(covering[0], formula, cells):
                return
            self._begin_change()
            ws.data_validations.dataValidation = [v for v in existing if v not in covering]
            # showDropDown is left alone: in the file format True means "hide the arrow".
            validation = DataValidation(type="list", formula1=formula, allow_blank=True)
            validation.add(cells)
            ws.add_data_validation(validation)

    # ---- Backups ---------------------------------------------------------------------

    def backup_files(self) -> list[Path]:
        """Existing automatic backups of this workbook, oldest first. (The copy `restore` keeps
        of the file it replaces has a different name, so it is never pruned or restored.)"""
        pattern = re.compile(re.escape(self.path.stem) + r"-\d{8}-\d{6}\.xlsx")
        if not self.backups_dir.is_dir():
            return []
        return sorted(p for p in self.backups_dir.iterdir() if pattern.fullmatch(p.name))

    def backup(self, now: datetime | None = None) -> Path:
        """Copy the workbook into the backups folder and keep only the newest MAX_BACKUPS.
        Copying works while Excel has the file open."""
        if not self.exists():
            raise WorkbookMissing(f"There is no job list at {self.path} to back up.")
        stamp = (now or datetime.now()).strftime(BACKUP_STAMP_FORMAT)
        target = self.backups_dir / f"{self.path.stem}-{stamp}.xlsx"
        try:
            self.backups_dir.mkdir(parents=True, exist_ok=True)
            partial = target.with_name(f".{target.name}.partial")
            shutil.copyfile(self.path, partial)
            os.replace(partial, target)
            for old in self.backup_files()[:-MAX_BACKUPS]:
                old.unlink(missing_ok=True)
        except OSError as exc:
            raise WorkbookError(
                f"Couldn't make a backup copy of your job list in {self.backups_dir}: {exc}"
            ) from exc
        return target

    def restore_latest_backup(self, now: datetime | None = None) -> tuple[Path, Path | None]:
        """Put the newest backup in place of the workbook. The file it replaces is first copied
        into the backups folder as "<name>-before-restore-<time>.xlsx". Returns (the backup
        that was restored, where the replaced file went or None if there was none)."""
        backups = self.backup_files()
        if not backups:
            raise WorkbookError(f"There are no backups of your job list in {self.backups_dir}.")
        self.require_closed()
        self._make_folder()
        stamp = (now or datetime.now()).strftime(BACKUP_STAMP_FORMAT)
        with self._exclusive_access():
            kept = self._keep_replaced_file(stamp)
            temp = self._new_temp_file()
            try:
                shutil.copyfile(backups[-1], temp)
            except OSError as exc:
                temp.unlink(missing_ok=True)
                raise WorkbookError(f"Couldn't restore {backups[-1].name}: {exc}") from exc
            self._install(temp)
        return backups[-1], kept

    def _keep_replaced_file(self, stamp: str) -> Path | None:
        if not self.exists():
            return None
        kept = self.backups_dir / f"{self.path.stem}-before-restore-{stamp}.xlsx"
        try:
            self.backups_dir.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(self.path, kept)
        except OSError as exc:
            raise WorkbookError(
                f"Couldn't keep a copy of the current file before restoring: {exc}. "
                "Nothing was changed."
            ) from exc
        return kept

    # ---- Opening it for the person ------------------------------------------------------

    def open_in_app(self) -> None:
        """Open the file in whatever program the computer uses for .xlsx files."""
        if not self.exists():
            raise WorkbookMissing(f"There is no job list at {self.path} to open.")
        try:
            if sys.platform == "win32":
                os.startfile(self.path)
            else:
                opener = "open" if sys.platform == "darwin" else "xdg-open"
                subprocess.run([opener, str(self.path)], check=True)
        except (OSError, subprocess.CalledProcessError) as exc:
            raise WorkbookError(
                f"Couldn't open {self.path}. Your computer needs a program that opens Excel "
                "files: install Microsoft Excel, or the free LibreOffice, then open the file "
                "yourself."
            ) from exc
